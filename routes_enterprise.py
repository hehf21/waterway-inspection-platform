# -*- coding: utf-8 -*-
"""路由模块：企业名录 / 一企一档 / 船舶清单 / 批量导入导出（由 app.py 拆分生成）"""
from fastapi import APIRouter, File, Form, Request, UploadFile

from webcore import *  # noqa: F401,F403

router = APIRouter()

# ---------------- 企业管理 ----------------
@router.get("/enterprises", response_class=HTMLResponse)
def ent_list(request: Request, q: str = "", p: int = 1):
    user, err = require(request, GOV_ROLES + ("sysadmin",))
    if err: return err
    conn = get_db()
    page_size = 50
    p = max(1, p)
    where, args = "", []
    if q:
        where = " WHERE name LIKE ? OR credit_code LIKE ?"
        args = [f"%{q}%", f"%{q}%"]
    total = conn.execute("SELECT COUNT(*) c FROM enterprises" + where, args).fetchone()["c"]
    rows = conn.execute("SELECT * FROM enterprises" + where + " ORDER BY id LIMIT ? OFFSET ?",
                        args + [page_size, (p - 1) * page_size]).fetchall()
    conn.close()
    pages = max(1, (total + page_size - 1) // page_size)
    return render(request, "enterprises.html",
                  {"rows": rows, "q": q, "p": p, "pages": pages, "total": total}, user)


@router.post("/enterprises/save")
async def ent_save(request: Request, eid: int = Form(0), name: str = Form(...), credit_code: str = Form(""),
                   license_no: str = Form(""), license_type: str = Form(""), legal_person: str = Form(""),
                   contact: str = Form(""), phone: str = Form(""), address: str = Form(""), scope: str = Form(""),
                   remark: str = Form(""), csrf: str = Form("")):
    user, err = require(request, ("gov_admin", "sysadmin"))
    if err: return err
    if bad_csrf(user, csrf):
        return HTMLResponse("表单已过期，请返回刷新后重试", status_code=403)
    conn = get_db()
    # 名称/信用代码去重（与批量导入口径一致，审计 P2-6）
    dup = conn.execute("SELECT id FROM enterprises WHERE name=? OR (?<>'' AND credit_code=?)",
                       (name, credit_code, credit_code)).fetchone()
    if dup and dup["id"] != (eid or 0):
        conn.close()
        return RedirectResponse("/enterprises?msg=保存失败：企业名称或信用代码已存在（若确为同一企业请直接编辑原记录）",
                                status_code=302)
    vals = (name, credit_code, license_no, license_type, legal_person, contact, phone, address, scope, remark)
    if eid:
        old = conn.execute("SELECT * FROM enterprises WHERE id=?", (eid,)).fetchone()
        if not old:
            conn.close(); return HTMLResponse("企业不存在", status_code=404)
        conn.execute("UPDATE enterprises SET name=?,credit_code=?,license_no=?,license_type=?,legal_person=?,"
                     "contact=?,phone=?,address=?,scope=?,remark=? WHERE id=?", vals + (eid,))
        fields = ["name", "credit_code", "license_no", "license_type", "legal_person", "contact", "phone",
                  "address", "scope", "remark"]
        audit(conn, user["username"], "修改企业", "enterprise", eid, name,
              before=json.dumps({f: old[f] for f in fields}, ensure_ascii=False),
              after=json.dumps(dict(zip(fields, vals)), ensure_ascii=False))
    else:
        cur = conn.execute("INSERT INTO enterprises(name,credit_code,license_no,license_type,legal_person,"
                           "contact,phone,address,scope,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                           vals + (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),))
        audit(conn, user["username"], "新增企业", "enterprise", cur.lastrowid, name)
    conn.commit(); conn.close()
    return RedirectResponse("/enterprises?msg=保存成功", status_code=302)


@router.get("/enterprises/{eid}/profile", response_class=HTMLResponse)
def ent_profile(request: Request, eid: int):
    """一企一档：企业基本信息 + 历次检查 + 问题台账 + 整改率"""
    user, err = require(request, GOV_ROLES + ("sysadmin", "enterprise"))
    if err: return err
    conn = get_db()
    ent = conn.execute("SELECT * FROM enterprises WHERE id=?", (eid,)).fetchone()
    if not ent:
        conn.close(); return HTMLResponse("企业不存在", status_code=404)
    if user["role"] == "enterprise" and user["enterprise_id"] != eid:
        conn.close(); return HTMLResponse("无权查看", status_code=403)
    ins_rows = conn.execute("SELECT * FROM inspections WHERE enterprise_id=? ORDER BY id DESC", (eid,)).fetchall()
    prob_rows = conn.execute(
        "SELECT p.*, i.code ins_code, i.check_date FROM problems p JOIN inspections i ON i.id=p.inspection_id"
        " WHERE i.enterprise_id=? ORDER BY p.id DESC", (eid,)).fetchall()
    today = date.today().isoformat()
    p_total = len(prob_rows)
    p_pass = sum(1 for p in prob_rows if p["status"] == "passed")
    p_over = sum(1 for p in prob_rows if p["status"] in ("pending", "returned") and p["deadline"] and p["deadline"] < today)
    probs = []
    for p in prob_rows:
        d = dict(p)
        d["seq_label"] = f"{p['ins_code']}-P{p['seq']}"
        probs.append(d)
    ships = conn.execute("SELECT * FROM ships WHERE enterprise_id=? ORDER BY id", (eid,)).fetchall()
    conn.close()
    return render(request, "ent_profile.html",
                  {"ent": ent, "ins_rows": ins_rows, "probs": probs, "ships": ships,
                   "p_total": p_total, "p_pass": p_pass, "p_over": p_over,
                   "rate": f"{p_pass / p_total * 100:.1f}%" if p_total else "—"}, user)


@router.post("/enterprises/self")
async def ent_self(request: Request, csrf: str = Form(""), contact: str = Form(""),
                   phone: str = Form(""), address: str = Form("")):
    """企业用户维护本企业联系信息（仅限安全联系人/电话/地址，审计 P2-9）"""
    user, err = require(request, ("enterprise",))
    if err: return err
    if bad_csrf(user, csrf):
        return HTMLResponse("表单已过期，请返回刷新后重试", status_code=403)
    eid = user["enterprise_id"]
    conn = get_db()
    old = conn.execute("SELECT * FROM enterprises WHERE id=?", (eid,)).fetchone()
    if not old:
        conn.close(); return HTMLResponse("企业不存在", status_code=404)
    conn.execute("UPDATE enterprises SET contact=?, phone=?, address=? WHERE id=?",
                 (contact.strip(), phone.strip(), address.strip(), eid))
    audit(conn, user["username"], "企业自助维护信息", "enterprise", eid, old["name"],
          before=json.dumps({"contact": old["contact"], "phone": old["phone"],
                             "address": old["address"]}, ensure_ascii=False),
          after=json.dumps({"contact": contact.strip(), "phone": phone.strip(),
                            "address": address.strip()}, ensure_ascii=False))
    conn.commit(); conn.close()
    return RedirectResponse(f"/enterprises/{eid}/profile?msg=企业联系信息已更新", status_code=302)


@router.post("/enterprises/{eid}/ships/save")
async def ship_save(request: Request, eid: int, sid: int = Form(0), name: str = Form(...),
                    ship_no: str = Form(""), ship_type: str = Form(""), dwt: str = Form(""),
                    gt: str = Form(""), built_date: str = Form(""), license_no: str = Form(""),
                    cert_status: str = Form(""), remark: str = Form(""), csrf: str = Form("")):
    """船舶清单维护：企业维护本企业船舶；gov_admin/sysadmin 可维护任意企业（审计 P2-9）"""
    user, err = require(request, ("enterprise", "gov_admin", "sysadmin"))
    if err: return err
    if bad_csrf(user, csrf):
        return HTMLResponse("表单已过期，请返回刷新后重试", status_code=403)
    if user["role"] == "enterprise" and user["enterprise_id"] != eid:
        return HTMLResponse("无权操作", status_code=403)
    fields = ["name", "ship_no", "ship_type", "dwt", "gt", "built_date", "license_no", "cert_status", "remark"]
    vals = tuple(v.strip() for v in (name, ship_no, ship_type, dwt, gt, built_date, license_no, cert_status, remark))
    conn = get_db()
    if sid:
        old = conn.execute("SELECT * FROM ships WHERE id=? AND enterprise_id=?", (sid, eid)).fetchone()
        if not old:
            conn.close(); return HTMLResponse("船舶不存在", status_code=404)
        conn.execute("UPDATE ships SET name=?,ship_no=?,ship_type=?,dwt=?,gt=?,built_date=?,license_no=?,"
                     "cert_status=?,remark=? WHERE id=?", vals + (sid,))
        audit(conn, user["username"], "修改船舶", "ship", sid, name,
              before=json.dumps({f: old[f] for f in fields}, ensure_ascii=False),
              after=json.dumps(dict(zip(fields, vals)), ensure_ascii=False))
    else:
        cur = conn.execute("INSERT INTO ships(enterprise_id,name,ship_no,ship_type,dwt,gt,built_date,license_no,"
                           "cert_status,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                           (eid,) + vals + (datetime.now().strftime("%Y-%m-%d %H:%M:%S"),))
        audit(conn, user["username"], "新增船舶", "ship", cur.lastrowid, name)
    conn.commit(); conn.close()
    return RedirectResponse(f"/enterprises/{eid}/profile?msg=船舶信息已保存", status_code=302)


@router.post("/ships/{sid}/delete")
def ship_delete(request: Request, sid: int, csrf: str = Form("")):
    user, err = require(request, ("enterprise", "gov_admin", "sysadmin"))
    if err: return err
    if bad_csrf(user, csrf):
        return HTMLResponse("表单已过期，请返回刷新后重试", status_code=403)
    conn = get_db()
    ship = conn.execute("SELECT * FROM ships WHERE id=?", (sid,)).fetchone()
    if not ship:
        conn.close(); return HTMLResponse("船舶不存在", status_code=404)
    if user["role"] == "enterprise" and user["enterprise_id"] != ship["enterprise_id"]:
        conn.close(); return HTMLResponse("无权操作", status_code=403)
    conn.execute("DELETE FROM ships WHERE id=?", (sid,))
    audit(conn, user["username"], "删除船舶", "ship", sid, ship["name"])
    conn.commit(); conn.close()
    return RedirectResponse(f"/enterprises/{ship['enterprise_id']}/profile?msg=船舶已删除", status_code=302)


# ---------------- 企业批量导入 / 检查清单导出 ----------------
@router.get("/enterprises/template")
def ent_template(request: Request):
    user, err = require(request, ("gov_admin", "sysadmin"))
    if err: return err
    import exportgen
    conn = get_db()
    audit(conn, user["username"], "下载企业导入模板", "enterprise", "")
    conn.commit(); conn.close()
    return Response(exportgen.ent_template_xlsx(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers=dl_header("企业名录批量导入模板.xlsx"))


@router.post("/enterprises/import")
async def ent_import(request: Request, csrf: str = Form(""), file: UploadFile = File(...)):
    user, err = require(request, ("gov_admin", "sysadmin"))
    if err: return err
    if bad_csrf(user, csrf):
        return HTMLResponse("表单已过期，请返回刷新后重试", status_code=403)
    from openpyxl import load_workbook
    # 限读 10MB：超限直接拒绝，不把大文件整包读进内存（与上传同口径，审计 P2-4）
    max_import = 10 * 1024 * 1024
    raw = file.file.read(max_import + 1)
    if len(raw) > max_import:
        return HTMLResponse("导入文件超过 10MB 上限，请拆分后分批导入", status_code=400)
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True)
    except Exception:
        return HTMLResponse("文件无法解析，请使用 .xlsx 模板", status_code=400)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return RedirectResponse("/enterprises?msg=表格为空", status_code=302)
    if len(rows) - 1 > 500:
        return HTMLResponse("单次最多导入 500 家（当前 %d 行），请拆分后分批导入" % (len(rows) - 1), status_code=400)
    head = [str(h or "").strip() for h in rows[0]]
    alias = {"企业名称": ["企业名称", "名称", "单位名称"], "统一社会信用代码": ["统一社会信用代码", "信用代码"],
             "经营许可证号": ["经营许可证号", "许可证号"], "经营类别": ["经营类别", "类别"],
             "法定代表人": ["法定代表人", "法人"], "安全联系人": ["安全联系人", "联系人"],
             "联系电话": ["联系电话", "电话"], "地址": ["地址"], "核定经营范围": ["核定经营范围", "经营范围"]}
    idx = {}
    for field, names in alias.items():
        for i, h in enumerate(head):
            if h in names:
                idx[field] = i
                break
    if "企业名称" not in idx:
        return HTMLResponse("表头缺少“企业名称”列，请使用导入模板", status_code=400)

    def cell(row, field):
        i = idx.get(field)
        return str(row[i]).strip() if i is not None and i < len(row) and row[i] is not None else ""

    conn = get_db()
    added = skipped = 0
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for row in rows[1:]:
        name = cell(row, "企业名称")
        if not name:
            continue
        code_no = cell(row, "统一社会信用代码")
        dup = conn.execute("SELECT 1 FROM enterprises WHERE name=? OR (?<>'' AND credit_code=?)",
                           (name, code_no, code_no)).fetchone()
        if dup:
            skipped += 1
            continue
        conn.execute("INSERT INTO enterprises(name,credit_code,license_no,license_type,legal_person,contact,"
                     "phone,address,scope,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                     (name, code_no, cell(row, "经营许可证号"), cell(row, "经营类别"), cell(row, "法定代表人"),
                      cell(row, "安全联系人"), cell(row, "联系电话"), cell(row, "地址"),
                      cell(row, "核定经营范围"), "批量导入", now))
        added += 1
    audit(conn, user["username"], "批量导入企业", "enterprise", "", f"新增{added}家，跳过{skipped}家")
    conn.commit(); conn.close()
    return RedirectResponse(f"/enterprises?msg=导入完成：新增{added}家，跳过重复{skipped}家", status_code=302)


@router.get("/items/export")
def items_export(request: Request, fmt: str = "xlsx"):
    user, err = require(request, GOV_ROLES + ("sysadmin",))
    if err: return err
    import exportgen
    conn = get_db()
    rows = [dict(r) for r in conn.execute("SELECT * FROM check_items WHERE active=1 ORDER BY code")]
    audit(conn, user["username"], "导出检查项目库", "check_item", "", f"fmt={fmt} 共{len(rows)}项")
    conn.commit(); conn.close()
    if fmt == "docx":
        data = exportgen.items_docx(rows, "水路运输企业检查内容清单（检查项目库）")
        return Response(data,
                        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers=dl_header("检查内容清单.docx"))
    data = exportgen.items_xlsx(rows)
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers=dl_header("检查项目库.xlsx"))


@router.get("/templates/{tid}/export")
def tpl_export(request: Request, tid: int, fmt: str = "xlsx"):
    user, err = require(request, GOV_ROLES + ("sysadmin",))
    if err: return err
    import exportgen
    conn = get_db()
    tpl = conn.execute("SELECT * FROM check_templates WHERE id=?", (tid,)).fetchone()
    if not tpl:
        conn.close(); return HTMLResponse("检查表不存在", status_code=404)
    rows = [dict(r) for r in conn.execute(
        "SELECT ci.* FROM template_items ti JOIN check_items ci ON ci.id=ti.item_id"
        " WHERE ti.template_id=? ORDER BY ti.sort_order", (tid,))]
    audit(conn, user["username"], "导出检查表", "template", tid, f"{tpl['name']} v{tpl['version']} fmt={fmt}")
    conn.commit(); conn.close()
    sub = f"{tpl['name']}（v{tpl['version']}）　{tpl['description']}"
    if fmt == "docx":
        data = exportgen.items_docx(rows, tpl["name"], sub)
        return Response(data,
                        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        headers=dl_header(f"{tpl['name']}.docx"))
    data = exportgen.items_xlsx(rows, tpl["name"])
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers=dl_header(f"{tpl['name']}.xlsx"))


