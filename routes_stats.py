# -*- coding: utf-8 -*-
"""路由模块：督办台账 / 统计分析 / 月度考核汇总（由 app.py 拆分生成）"""
from fastapi import APIRouter, File, Form, Request, UploadFile

from webcore import *  # noqa: F401,F403

router = APIRouter()

# ---------------- 督办台账 ----------------
@router.get("/overdue", response_class=HTMLResponse)
def overdue_list(request: Request):
    user, err = require(request, GOV_ROLES + ("enterprise", "sysadmin"))
    if err: return err
    conn = get_db()
    today_d = date.today()
    today = today_d.isoformat()
    soon_day = (today_d + timedelta(days=7)).isoformat()
    scope = "" if user["role"] != "enterprise" else " AND i.enterprise_id=?"

    def q(deadline_sql, args):
        sql = ("SELECT p.*, i.code ins_code, i.enterprise_id, e.name ent_name FROM problems p"
               " JOIN inspections i ON i.id=p.inspection_id LEFT JOIN enterprises e ON e.id=i.enterprise_id"
               " WHERE p.status IN ('pending','returned') AND p.deadline<>'' " + deadline_sql + scope +
               " ORDER BY p.deadline")
        return conn.execute(sql, args).fetchall()

    base_args = [] if user["role"] != "enterprise" else [user["enterprise_id"]]
    rows = []
    for r in q("AND p.deadline<?", [today] + base_args):
        d = dict(r)
        d["days_over"] = (today_d - date.fromisoformat(r["deadline"])).days
        rows.append(d)
    soon = []
    for r in q("AND p.deadline>=? AND p.deadline<=?", [today, soon_day] + base_args):
        d = dict(r)
        d["days_left"] = (date.fromisoformat(r["deadline"]) - today_d).days
        soon.append(d)
    conn.close()
    return render(request, "overdue.html", {"rows": rows, "soon": soon, "today": today}, user)


# ---------------- 批量打印检查表 / 月度考核汇总表 ----------------
@router.get("/print-batch", response_class=HTMLResponse)
def ins_print_batch(request: Request, ids: str = ""):
    """批量打印检查表：多条合成一份 A4 文档（每条从新页起），浏览器一次打印"""
    user, err = require(request)
    if err: return err
    id_list = [int(x) for x in ids.split(",") if x.strip().isdigit()]
    if not id_list:
        return HTMLResponse("请先在检查记录列表勾选记录，再点“批量打印检查表”", status_code=400)
    conn = get_db()
    records = []
    for iid in id_list[:50]:
        ins = conn.execute("SELECT i.*, e.name ent_name, e.credit_code, e.license_no FROM inspections i"
                           " LEFT JOIN enterprises e ON e.id=i.enterprise_id WHERE i.id=?", (iid,)).fetchone()
        if not ins:
            continue
        if user["role"] not in GOV_ROLES + ("sysadmin",) and ins["enterprise_id"] != user["enterprise_id"]:
            continue
        if user["role"] == "enterprise" and ins["status"] == "draft":
            continue
        items = [dict(r) for r in conn.execute(
            "SELECT * FROM inspection_items WHERE inspection_id=? ORDER BY sort_order", (iid,))]
        probs = [dict(r) for r in conn.execute("SELECT * FROM problems WHERE inspection_id=? ORDER BY seq", (iid,))]
        signs = {s["sign_type"]: s for s in
                 conn.execute("SELECT * FROM signatures WHERE inspection_id=?", (iid,)).fetchall()}
        d = dict(ins)
        d["status_cn"] = REC_STATUS_CN.get(ins["status"], ins["status"])
        records.append({"ins": d, "ins_items": items, "probs": probs, "signs": signs})
    audit(conn, user["username"], "批量打印检查表", "inspection", ",".join(map(str, id_list)),
          f"选中{len(id_list)}条 输出{len(records)}条")
    conn.commit(); conn.close()
    return render(request, "inspection_print_batch.html", {"records": records, "org": ORG_NAME}, user)


@router.get("/stats/monthly", response_class=HTMLResponse)
def stats_monthly(request: Request, month: str = ""):
    """月度考核汇总表（可打印）：检查量/问题整改/计划执行/双随机批次"""
    user, err = require(request, GOV_ROLES + ("sysadmin",))
    if err: return err
    month = month if re.match(r"^\d{4}-\d{2}$", month or "") else date.today().strftime("%Y-%m")
    y, m = int(month[:4]), int(month[5:7])
    m0 = f"{month}-01"
    m1 = date(y + (m == 12), (m % 12) + 1, 1).isoformat()
    conn = get_db()
    head = conn.execute(
        "SELECT COUNT(*) ins_cnt, COUNT(DISTINCT enterprise_id) ent_cnt,"
        " SUM(CASE WHEN conclusion='未发现问题' THEN 1 ELSE 0 END) clean_cnt,"
        " SUM(CASE WHEN conclusion='责发整改' THEN 1 ELSE 0 END) rect_cnt,"
        " SUM(CASE WHEN conclusion='移送线索' THEN 1 ELSE 0 END) move_cnt"
        " FROM inspections WHERE check_date>=? AND check_date<? AND status<>'draft'", (m0, m1)).fetchone()
    prob = conn.execute(
        "SELECT COUNT(*) p_cnt, SUM(CASE WHEN p.status='passed' THEN 1 ELSE 0 END) passed,"
        " SUM(CASE WHEN p.status IN ('pending','returned') AND p.deadline<>'' AND p.deadline<? THEN 1 ELSE 0 END) overdue"
        " FROM problems p JOIN inspections i ON i.id=p.inspection_id WHERE i.check_date>=? AND i.check_date<?",
        (date.today().isoformat(), m0, m1)).fetchone()
    by_type = conn.execute("SELECT check_type, COUNT(*) c FROM inspections WHERE check_date>=? AND check_date<?"
                           " AND status<>'draft' GROUP BY check_type", (m0, m1)).fetchall()
    by_cat = conn.execute("SELECT ci.category, COUNT(*) c FROM problems p JOIN check_items ci ON ci.id=p.item_id"
                          " JOIN inspections i ON i.id=p.inspection_id WHERE i.check_date>=? AND i.check_date<?"
                          " GROUP BY ci.category ORDER BY c DESC", (m0, m1)).fetchall()
    draws = conn.execute("SELECT COUNT(*) c FROM random_draws WHERE created_at>=? AND created_at<?",
                         (m0 + " 00:00:00", m1 + " 00:00:00")).fetchone()["c"]
    plan_total = conn.execute("SELECT COALESCE(SUM(count_plan),0) c FROM plan_items").fetchone()["c"]
    plan_done_m = conn.execute("SELECT COUNT(*) c FROM inspections WHERE plan_item_id>0"
                               " AND check_date>=? AND check_date<? AND status<>'draft'", (m0, m1)).fetchone()["c"]
    plan_done_all = conn.execute("SELECT COUNT(*) c FROM inspections WHERE plan_item_id>0 AND status<>'draft'"
                                 ).fetchone()["c"]
    conn.close()
    ctx = {"month": month, "head": head, "prob": prob, "by_type": by_type, "by_cat": by_cat,
           "draws": draws, "plan_total": plan_total, "plan_done_m": plan_done_m,
           "plan_done_all": plan_done_all, "org": ORG_NAME}
    return render(request, "monthly.html", ctx, user)


REPORT_TPL = os.path.join(BASE, "data", "report_template.xlsx")


def _fill_report_tpl(path, vals):
    """自定义考核报表模板套打：把单元格里的 {{key}} 替换为当月数据（文本替换，表格行不扩展）"""
    from openpyxl import load_workbook
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    s = cell.value
                    for k, v in vals.items():
                        s = s.replace("{{" + k + "}}", str(v))
                    cell.value = s
    return wb


@router.post("/stats/monthly/template")
def stats_monthly_template(request: Request, csrf: str = Form(""), file: UploadFile = File(...)):
    """上传考核报表模板（.xlsx，单元格写 {{month}}/{{ins_cnt}} 等占位符），导出时自动套打"""
    user, err = require(request, ("gov_admin",))
    if err: return err
    if bad_csrf(user, csrf):
        return HTMLResponse("表单已过期，请返回刷新后重试", status_code=403)
    raw = file.file.read(2 * 1024 * 1024)
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(raw))
        assert wb.worksheets
    except Exception:
        return HTMLResponse("模板无法解析，请上传 .xlsx 格式", status_code=400)
    os.makedirs(os.path.dirname(REPORT_TPL), exist_ok=True)
    with open(REPORT_TPL, "wb") as fp:
        fp.write(raw)
    conn = get_db()
    audit(conn, user["username"], "上传考核报表模板", "stats", "", file.filename or "")
    conn.commit(); conn.close()
    return RedirectResponse("/stats/monthly?msg=考核报表模板已保存，导出将按模板套打", status_code=302)


@router.get("/stats/monthly/export")
def stats_monthly_export(request: Request, month: str = ""):
    user, err = require(request, GOV_ROLES + ("sysadmin",))
    if err: return err
    month = month if re.match(r"^\d{4}-\d{2}$", month or "") else date.today().strftime("%Y-%m")
    y, m = int(month[:4]), int(month[5:7])
    m0 = f"{month}-01"
    m1 = date(y + (m == 12), (m % 12) + 1, 1).isoformat()
    conn = get_db()
    wb = Workbook()
    ws = wb.active
    ws.title = "月度汇总"
    head = conn.execute(
        "SELECT COUNT(*) ins_cnt, COUNT(DISTINCT enterprise_id) ent_cnt,"
        " SUM(CASE WHEN conclusion='未发现问题' THEN 1 ELSE 0 END) clean_cnt,"
        " SUM(CASE WHEN conclusion='责发整改' THEN 1 ELSE 0 END) rect_cnt,"
        " SUM(CASE WHEN conclusion='移送线索' THEN 1 ELSE 0 END) move_cnt"
        " FROM inspections WHERE check_date>=? AND check_date<? AND status<>'draft'", (m0, m1)).fetchone()
    prob = conn.execute(
        "SELECT COUNT(*) p_cnt, SUM(CASE WHEN p.status='passed' THEN 1 ELSE 0 END) passed,"
        " SUM(CASE WHEN p.status IN ('pending','returned') AND p.deadline<>'' AND p.deadline<? THEN 1 ELSE 0 END) overdue"
        " FROM problems p JOIN inspections i ON i.id=p.inspection_id WHERE i.check_date>=? AND i.check_date<?",
        (date.today().isoformat(), m0, m1)).fetchone()
    ws.append(["月份", month])
    ws.append(["检查次数", head["ins_cnt"] or 0])
    ws.append(["涉及企业数", head["ent_cnt"] or 0])
    ws.append(["结论-未发现问题", head["clean_cnt"] or 0])
    ws.append(["结论-责发整改", head["rect_cnt"] or 0])
    ws.append(["结论-移送线索", head["move_cnt"] or 0])
    ws.append(["发现问题数", prob["p_cnt"] or 0])
    ws.append(["已销号", prob["passed"] or 0])
    rate = f"{(prob['passed'] or 0) / prob['p_cnt'] * 100:.1f}%" if prob["p_cnt"] else "-"
    ws.append(["整改完成率", rate])
    ws.append(["其中逾期未整改", prob["overdue"] or 0])
    vals = {"month": month, "org": ORG_NAME, "ins_cnt": head["ins_cnt"] or 0, "ent_cnt": head["ent_cnt"] or 0,
            "clean_cnt": head["clean_cnt"] or 0, "rect_cnt": head["rect_cnt"] or 0,
            "move_cnt": head["move_cnt"] or 0, "p_cnt": prob["p_cnt"] or 0,
            "passed": prob["passed"] or 0, "rate": rate, "overdue": prob["overdue"] or 0}
    ws2 = wb.create_sheet("按检查类型")
    ws2.append(["检查类型", "次数"])
    for r2 in conn.execute("SELECT check_type, COUNT(*) c FROM inspections WHERE check_date>=? AND check_date<?"
                           " AND status<>'draft' GROUP BY check_type", (m0, m1)):
        ws2.append([r2["check_type"], r2["c"]])
    ws3 = wb.create_sheet("按问题类别")
    ws3.append(["检查项类别", "问题数"])
    for r2 in conn.execute("SELECT ci.category, COUNT(*) c FROM problems p JOIN check_items ci ON ci.id=p.item_id"
                           " JOIN inspections i ON i.id=p.inspection_id WHERE i.check_date>=? AND i.check_date<?"
                           " GROUP BY ci.category ORDER BY c DESC", (m0, m1)):
        ws3.append([r2["category"], r2["c"]])
    audit(conn, user["username"], "导出月度考核汇总表", "stats", "", month)
    conn.commit(); conn.close()
    if os.path.exists(REPORT_TPL):   # 自定义模板套打（未上传模板则用默认格式）
        wb = _fill_report_tpl(REPORT_TPL, vals)
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers=dl_header(f"月度考核汇总表_{month}.xlsx"))


# ---------------- 统计分析 ----------------
@router.get("/stats", response_class=HTMLResponse)
def stats_page(request: Request, date_from: str = "", date_to: str = ""):
    user, err = require(request, GOV_ROLES + ("sysadmin",))
    if err: return err
    conn = get_db()
    # 时间段筛选（审计 P2-11）：按检查日期过滤，留空统计全部
    filt, fargs = "", []
    if date_from:
        filt += " AND check_date>=?"; fargs.append(date_from)
    if date_to:
        filt += " AND check_date<=?"; fargs.append(date_to)
    by_ent = conn.execute(
        "SELECT e.name, COUNT(DISTINCT i.id) ins_cnt, COUNT(DISTINCT p.id) prob_cnt,"
        " COUNT(DISTINCT CASE WHEN p.status='passed' THEN p.id END) pass_cnt"
        " FROM enterprises e LEFT JOIN inspections i ON i.enterprise_id=e.id" + filt +
        " LEFT JOIN problems p ON p.inspection_id=i.id"
        " GROUP BY e.id ORDER BY ins_cnt DESC", fargs).fetchall()
    by_type = conn.execute("SELECT check_type, COUNT(*) c FROM inspections WHERE 1=1" + filt +
                           " GROUP BY check_type", fargs).fetchall()
    by_cat = conn.execute(
        "SELECT ci.category, COUNT(*) c FROM problems p JOIN check_items ci ON ci.id=p.item_id"
        " JOIN inspections i ON i.id=p.inspection_id WHERE 1=1" + filt +
        " GROUP BY ci.category ORDER BY c DESC", fargs).fetchall()
    conn.close()
    return render(request, "stats.html",
                  {"by_ent": by_ent, "by_type": by_type, "by_cat": by_cat,
                   "d_from": date_from, "d_to": date_to}, user)


@router.get("/stats/export")
def stats_export(request: Request, date_from: str = "", date_to: str = ""):
    user, err = require(request, GOV_ROLES + ("sysadmin",))
    if err: return err
    conn = get_db()
    filt, fargs = "", []
    if date_from:
        filt += " AND check_date>=?"; fargs.append(date_from)
    if date_to:
        filt += " AND check_date<=?"; fargs.append(date_to)
    wb = Workbook()
    ws = wb.active
    ws.title = "企业检查统计"
    ws.append(["企业名称", "检查次数", "发现问题数", "已整改数", "整改完成率"])
    for r in conn.execute(
            "SELECT e.id, e.name, COUNT(DISTINCT i.id) ins_cnt, COUNT(DISTINCT p.id) prob_cnt,"
            " COUNT(DISTINCT CASE WHEN p.status='passed' THEN p.id END) pass_cnt"
            " FROM enterprises e LEFT JOIN inspections i ON i.enterprise_id=e.id" + filt +
            " LEFT JOIN problems p ON p.inspection_id=i.id"
            " GROUP BY e.id", fargs):
        rate = f"{r['pass_cnt'] / r['prob_cnt'] * 100:.1f}%" if r["prob_cnt"] else "-"
        ws.append([r["name"], r["ins_cnt"], r["prob_cnt"], r["pass_cnt"], rate])
    ws2 = wb.create_sheet("检查记录明细")
    ws2.append(["检查编号", "企业", "检查日期", "类型", "结论", "状态", "检查人员"])
    for r in conn.execute("SELECT i.*, e.name ent_name FROM inspections i LEFT JOIN enterprises e ON e.id=i.enterprise_id"
                          " WHERE 1=1" + filt + " ORDER BY i.id DESC", fargs):
        ws2.append([r["code"], r["ent_name"], r["check_date"], r["check_type"], r["conclusion"],
                    REC_STATUS_CN.get(r["status"], r["status"]), r["inspectors"]])
    audit(conn, user["username"], "导出统计报表", "stats", "",
          f"企业检查统计+检查记录明细 区间{date_from or '不限'}~{date_to or '不限'}")
    conn.commit(); conn.close()
    buf = io.BytesIO()
    wb.save(buf)
    return Response(buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=stats.xlsx"})


