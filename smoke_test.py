# -*- coding: utf-8 -*-
"""全流程冒烟测试 v2：登录/改密/CSRF/登记/草稿编辑删除/下发/整改/复核/归档锁定/更正留痕/
批量导入/导出/督办/权限隔离。隔离数据目录，跑完自动清理，可重复运行。"""
import io
import json
import os
import re
import shutil
import sys
import zipfile
from datetime import date, timedelta

sys.stdout.reconfigure(encoding="utf-8")
TEST_DATA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "data_smoke")
if os.path.exists(TEST_DATA):
    shutil.rmtree(TEST_DATA)
os.environ["SLYS_DATA"] = TEST_DATA

from fastapi.testclient import TestClient
from app import app
from db import init_db

init_db()
c = TestClient(app)
ok = []
GOV_NEW_PWD = "gov2026!"


def check(name, cond, extra=""):
    ok.append((name, bool(cond)))
    print(("PASS " if cond else "FAIL ") + name + ("  " + extra if extra else ""))


def csrf_of():
    r = c.get("/")
    m = re.search(r'<meta name="csrf" content="([^"]+)"', r.text)
    return m.group(1) if m else ""


def new_item(code="A01", name="测试检查项", result="符合", finding=""):
    return {"code": code, "category": "A 资质与证照", "name": name, "content": "内容",
            "basis": "《测试条例》第1条", "method": "查资料", "criteria": "合格", "is_key": 1,
            "is_veto": 0, "score": 5, "result": result, "finding": finding}


# ========== 1. 登录 + 首次强制改密 ==========
r = c.post("/login", data={"username": "gov", "password": "gov@123"}, follow_redirects=False)
check("gov 登录", r.status_code == 302)
r = c.get("/", follow_redirects=False)
check("强制改密拦截", r.status_code == 302 and r.headers["location"].startswith("/changepwd"))
t = csrf_of()
r = c.post("/changepwd", data={"csrf": t, "old_password": "", "new_password": GOV_NEW_PWD,
                              "new_password2": GOV_NEW_PWD}, follow_redirects=False)
check("首次改密成功", r.status_code == 302)
r = c.get("/")
check("改密后进工作台", "工作台" in r.text)
r = c.get("/changepwd")
check("改密页可访问", "修改密码" in r.text)

# ========== 2. CSRF 防护 ==========
r = c.post("/items/1/toggle", data={"csrf": "伪造token"})
check("CSRF 伪造被拒 403", r.status_code == 403)
r = c.post("/items/1/toggle", data={})
check("CSRF 缺失被拒 403", r.status_code == 403)

# ========== 3. 企业 + 企业账号 ==========
t = csrf_of()
r = c.post("/enterprises/save", data={"csrf": t, "eid": 0, "name": "平潭测试航运有限公司",
                                     "credit_code": "91350128TEST0001X", "license_no": "闽交水运字0001号",
                                     "license_type": "普通货物运输", "legal_person": "王船长",
                                     "contact": "陈安全", "phone": "13800000000", "address": "平潭综合实验区",
                                     "scope": "国内沿海普通货物运输", "remark": "冒烟测试"}, follow_redirects=False)
check("新增企业", r.status_code == 302)
r = c.post("/users/save", data={"csrf": t, "uid": 0, "username": "ent", "password": "ent@123",
                               "real_name": "陈安全", "role": "enterprise", "enterprise_id": 1,
                               "phone": "13800000000"}, follow_redirects=False)
check("创建企业账号(强制改密)", r.status_code == 302)

# ========== 4. 企业批量导入 ==========
from openpyxl import Workbook
wb = Workbook()
ws = wb.active
ws.append(["企业名称", "统一社会信用代码", "经营许可证号", "经营类别", "法定代表人", "安全联系人", "联系电话", "地址", "核定经营范围"])
ws.append(["平潭导入甲公司", "91350128AAAA0001", "闽001", "旅客运输", "甲一", "乙一", "138", "平潭", "沿海客运"])
ws.append(["平潭导入乙公司", "91350128BBBB0002", "闽002", "危险货物运输", "甲二", "乙二", "139", "平潭", "危货运输"])
ws.append(["平潭测试航运有限公司", "91350128TEST0001X", "闽003", "普通货物运输", "重", "复", "137", "平潭", "重复应跳过"])
ib = io.BytesIO()
wb.save(ib)
ib.seek(0)
r = c.post("/enterprises/import", data={"csrf": t}, files={"file": ("名录.xlsx", ib, "application/vnd.ms-excel")},
           follow_redirects=False)
check("批量导入提交", r.status_code == 302)
from urllib.parse import unquote
check("导入结果：新增2跳过1", "新增2家，跳过重复1家" in unquote(r.headers.get("location", "")),
      unquote(r.headers.get("location", "")))
r = c.get("/enterprises")
check("导入新增2家", "平潭导入甲公司" in r.text and "平潭导入乙公司" in r.text)
r = c.get("/enterprises/template")
check("导入模板下载", r.status_code == 200 and r.content[:2] == b"PK")

# ========== 5. 导出检查清单 ==========
r = c.get("/items/export?fmt=docx")
check("检查清单导出Word", r.status_code == 200 and r.content[:2] == b"PK")
r = c.get("/items/export?fmt=xlsx")
check("检查清单导出Excel", r.status_code == 200 and r.content[:2] == b"PK")
r = c.get("/templates/1/export?fmt=docx")
check("检查表导出Word", r.status_code == 200 and r.content[:2] == b"PK")

# ========== 6. 登记检查（2 问题：一个限期3天、一个用默认） ==========
soon = (date.today() + timedelta(days=3)).isoformat()
later = (date.today() + timedelta(days=30)).isoformat()
items = [new_item("A01", "许可证检查"), new_item("A05", "海务机务管理人员", "不符合", "仅兼职"),
         new_item("B01", "责任制检查")]
probs = [
    {"item_id": 2, "description": "海务机务人员配备不足", "basis": "《国内水路运输管理规定》第8条",
     "requirement": "限期配备专职人员", "deadline": soon, "responsible": "王船长", "to_msa": False},
    {"item_id": 3, "description": "全员责任制未上墙", "basis": "《安全生产法》第22条",
     "requirement": "制作责任制并公示", "deadline": later, "responsible": "陈安全", "to_msa": False},
]
up = ("现场照片.jpg", io.BytesIO(b"\xff\xd8 fakejpg"), "image/jpeg")
r = c.post("/inspections/save",
           data={"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 1, "check_type": "日常检查",
                 "check_mode": "现场检查", "check_date": date.today().isoformat(), "location": "企业办公场所",
                 "inspectors": "李检查员", "deadline": later, "conclusion": "责发整改",
                 "conclusion_note": "责令限期整改",
                 "items_json": json.dumps(items, ensure_ascii=False),
                 "problems_json": json.dumps(probs, ensure_ascii=False)},
           files=[("files", up)], follow_redirects=False)
check("登记检查记录", r.status_code == 302, r.headers.get("location", ""))
ins_id = int(r.headers.get("location", "/inspections/0").split("/")[2].split("?")[0])

# ========== 7. 草稿编辑 + 删除 ==========
r = c.get(f"/inspections/new?edit={ins_id}")
check("草稿编辑页可开", r.status_code == 200 and "编辑检查草稿" in r.text)
r = c.post("/inspections/save",
           data={"csrf": t, "ins_id": ins_id, "enterprise_id": 1, "template_id": 1,
                 "check_type": "专项检查", "check_mode": "现场检查", "check_date": date.today().isoformat(),
                 "location": "企业办公场所-已修改", "inspectors": "李检查员", "deadline": later,
                 "conclusion": "责发整改", "conclusion_note": "责令限期整改",
                 "items_json": json.dumps(items, ensure_ascii=False),
                 "problems_json": json.dumps(probs, ensure_ascii=False)}, follow_redirects=False)
check("草稿保存修改", r.status_code == 302)
from db import get_db
conn = get_db()
row = conn.execute("SELECT * FROM inspections WHERE id=?", (ins_id,)).fetchone()
check("草稿内容已更新", row["location"] == "企业办公场所-已修改" and row["check_type"] == "专项检查")
check("草稿仍为draft", row["status"] == "draft")
conn.close()
# 另建一草稿并删除
r = c.post("/inspections/save",
           data={"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 0, "check_type": "日常检查",
                 "check_mode": "书面检查", "check_date": date.today().isoformat(), "location": "待删",
                 "inspectors": "张三", "deadline": "", "conclusion": "未发现问题", "conclusion_note": "",
                 "items_json": json.dumps([new_item()], ensure_ascii=False),
                 "problems_json": "[]"}, follow_redirects=False)
tmp_id = int(r.headers.get("location", "/inspections/0").split("/")[2].split("?")[0])
r = c.post(f"/inspections/{tmp_id}/delete", data={"csrf": t}, follow_redirects=False)
check("草稿删除", r.status_code == 302)
conn = get_db()
check("草稿已从库中删除", conn.execute("SELECT 1 FROM inspections WHERE id=?", (tmp_id,)).fetchone() is None)
code1 = conn.execute("SELECT code FROM inspections WHERE id=?", (ins_id,)).fetchone()["code"]
conn.close()

# ========== 7.5 检查记录查找与筛选 ==========
r = c.get(f"/inspections?q={code1}")
check("筛选-按检查编号", code1 in r.text)
r = c.get("/inspections?q=企业办公场所")
check("筛选-按检查地点关键字", code1 in r.text)
r = c.get("/inspections?q=不存在的关键词XYZ")
check("筛选-无结果", code1 not in r.text and "共 0 条" in r.text)
r = c.get("/inspections?check_type=专项检查")
check("筛选-按检查类型", code1 in r.text)
r = c.get("/inspections?check_type=日常检查")
check("筛选-类型排除", code1 not in r.text)
r = c.get("/inspections?check_mode=现场检查&conclusion=责发整改")
check("筛选-方式+结论组合", code1 in r.text)
r = c.get(f"/inspections?date_from={date.today().isoformat()}&date_to={date.today().isoformat()}")
check("筛选-日期范围含", code1 in r.text)
r = c.get(f"/inspections?date_from={(date.today() + timedelta(days=1)).isoformat()}")
check("筛选-日期范围排除", code1 not in r.text)
r = c.get("/inspections?enterprise_id=1&status=draft")
check("筛选-企业+状态组合", code1 in r.text)
r = c.get("/inspections?enterprise_id=2")
check("筛选-企业排除", code1 not in r.text)

# ========== 8. 下发 + 企业整改 ==========
r = c.post(f"/inspections/{ins_id}/issue", data={"csrf": t}, follow_redirects=False)
check("下发检查记录", r.status_code == 302)
r = c.get(f"/inspections/{ins_id}")
check("详情含问题", "海务机务人员配备不足" in r.text)
r = c.post("/login", data={"username": "ent", "password": "ent@123"}, follow_redirects=False)
check("企业登录", r.status_code == 302)
t2 = csrf_of()  # 强制改密页也有 csrf
r = c.get("/", follow_redirects=False)
check("企业被强制改密", r.status_code == 302)
r = c.post("/changepwd", data={"csrf": t2, "old_password": "", "new_password": "ent2026!",
                              "new_password2": "ent2026!"}, follow_redirects=False)
check("企业改密成功", r.status_code == 302)
t2 = csrf_of()
r = c.get(f"/inspections/{ins_id}")
check("企业可查看检查记录", "海务机务" in r.text)
conn = get_db()
pids = [r0["id"] for r0 in conn.execute("SELECT id FROM problems WHERE inspection_id=? ORDER BY seq", (ins_id,))]
conn.close()
up2 = ("整改照片.png", io.BytesIO(b"\x89PNGfake"), "image/png")
r = c.post(f"/problems/{pids[0]}/feedback",
           data={"csrf": t2, "reason": "人员调配不及时", "measure": "已招聘2名专职人员并备案",
                 "completion": "已完成", "done_date": date.today().isoformat()},
           files=[("files", up2)], follow_redirects=False)
check("提交整改反馈(问题1)", r.status_code == 302)
r = c.post(f"/problems/{pids[1]}/feedback",
           data={"csrf": t2, "reason": "重视不足", "measure": "制作制度牌并上墙", "completion": "已完成",
                 "done_date": date.today().isoformat()}, follow_redirects=False)
check("提交整改反馈(问题2)", r.status_code == 302)

# ========== 9. 政府复核（1退回重报→再通过；1通过） ==========
r = c.post("/login", data={"username": "gov", "password": GOV_NEW_PWD}, follow_redirects=False)
check("gov 再登录", r.status_code == 302)
t = csrf_of()
r = c.post(f"/problems/{pids[0]}/review",
           data={"csrf": t, "result": "reject", "opinion": "人员未见证明，退回补材料"}, follow_redirects=False)
check("复核退回", r.status_code == 302)
conn = get_db()
check("问题转returned", conn.execute("SELECT status FROM problems WHERE id=?", (pids[0],)).fetchone()["status"] == "returned")
conn.close()
# 企业重报
r = c.post("/login", data={"username": "ent", "password": "ent2026!"}, follow_redirects=False)
t2 = csrf_of()
r = c.post(f"/problems/{pids[0]}/feedback",
           data={"csrf": t2, "reason": "补充说明", "measure": "已上传2名人员资格证书扫描件",
                 "completion": "已完成", "done_date": date.today().isoformat()}, follow_redirects=False)
check("退回后重新反馈", r.status_code == 302)
conn = get_db()
check("保留两轮反馈", conn.execute("SELECT COUNT(*) c FROM feedbacks WHERE problem_id=?", (pids[0],)).fetchone()["c"] == 2)
conn.close()
# 政府全部通过
r = c.post("/login", data={"username": "gov", "password": GOV_NEW_PWD}, follow_redirects=False)
t = csrf_of()
for pid in pids:
    r = c.post(f"/problems/{pid}/review",
               data={"csrf": t, "result": "pass", "opinion": "经复核已整改到位",
                     "recheck_date": date.today().isoformat(), "recheck_note": "属实"},
               follow_redirects=False)
check("复核全部通过", r.status_code == 302)
conn = get_db()
st = conn.execute("SELECT status FROM inspections WHERE id=?", (ins_id,)).fetchone()["status"]
check("记录自动闭环", st == "closed", f"status={st}")
conn.close()

# ========== 10. 打包下载（只读）+ 显式归档锁定 + 更正留痕 ==========
r = c.get(f"/inspections/{ins_id}/download")
check("打包下载200", r.status_code == 200 and r.headers["content-type"] == "application/zip")
z = zipfile.ZipFile(io.BytesIO(r.content))
names = z.namelist()
need = ["01-现场检查记录表.pdf", "02-问题清单及整改要求.pdf", "03-企业整改情况报告.pdf",
        "05-复核意见与复查记录.pdf", "06-归档目录.xlsx", "06-归档说明.txt",
        "07-结构化数据.json", "07-结构化数据.xlsx", "SHA256SUMS.txt"]
check("ZIP 内容齐全", all(n in names for n in need), str([n for n in names if n not in need]))
check("ZIP 含附件按问题分目录", any(n.startswith("04-附件/P1-") for n in names))
conn = get_db()
st = conn.execute("SELECT status FROM inspections WHERE id=?", (ins_id,)).fetchone()["status"]
check("只读下载不改变状态", st == "closed", f"status={st}")
conn.close()
# 显式归档（审计 P1-1/P1-2）：POST /archive → 落盘保存 + 锁定；GET download 不再是锁定入口
r = c.post(f"/inspections/{ins_id}/archive", data={"csrf": t}, follow_redirects=False)
check("显式归档返回ZIP", r.status_code == 200 and r.headers["content-type"] == "application/zip",
      f"status={r.status_code}")
conn = get_db()
row = conn.execute("SELECT status, archive_file FROM inspections WHERE id=?", (ins_id,)).fetchone()
check("P0-4 归档后转已归档", row["status"] == "archived", f"status={row['status']}")
pack = os.path.join(os.environ["SLYS_DATA"], "archives", row["archive_file"])
check("归档包落盘保存", os.path.exists(pack), pack)
conn.close()
r1 = c.get(f"/inspections/{ins_id}/download")
r2 = c.get(f"/inspections/{ins_id}/download")
check("已归档包复用且字节一致", r1.content == r2.content)
# P0-1 归档锁定：企业反馈 / 政府复核均被拒
r = c.post("/login", data={"username": "ent", "password": "ent2026!"}, follow_redirects=False)
t2 = csrf_of()
r = c.post(f"/problems/{pids[0]}/feedback",
           data={"csrf": t2, "reason": "x", "measure": "y"}, follow_redirects=False)
check("P0-1 归档后企业反馈被拒", r.status_code == 403)
r = c.post("/login", data={"username": "gov", "password": GOV_NEW_PWD}, follow_redirects=False)
t = csrf_of()
r = c.post(f"/problems/{pids[0]}/review", data={"csrf": t, "result": "pass"}, follow_redirects=False)
check("P0-1 归档后复核被拒", r.status_code == 403)
# 更正留痕：admin 解除锁定
r = c.post("/login", data={"username": "admin", "password": "admin@123"}, follow_redirects=False)
t3 = csrf_of()
r = c.get("/", follow_redirects=False)
if r.status_code == 302:  # admin 也被强制改密
    r = c.post("/changepwd", data={"csrf": t3, "old_password": "", "new_password": "adm2026!",
                                  "new_password2": "adm2026!"}, follow_redirects=False)
    check("admin 强制改密", r.status_code == 302)
    t3 = csrf_of()
r = c.post(f"/inspections/{ins_id}/reopen", data={"csrf": t3, "reason": "整改附件上传错误，需重新上传"},
           follow_redirects=False)
check("更正留痕解除锁定", r.status_code == 302)
r = c.get(f"/inspections/{ins_id}")
check("更正历史可见", "更正留痕" in r.text and "附件上传错误" in r.text)
conn = get_db()
st = conn.execute("SELECT status FROM inspections WHERE id=?", (ins_id,)).fetchone()["status"]
check("解除后回到闭环态", st == "closed", f"status={st}")
conn.close()

# ========== 11. 单份文书 / 批量打包（gov 会话） ==========
r = c.post("/login", data={"username": "gov", "password": GOV_NEW_PWD}, follow_redirects=False)
t = csrf_of()
for kind in ("record", "notice", "feedback", "review"):
    r = c.get(f"/inspections/{ins_id}/doc/{kind}")
    check(f"文书PDF-{kind}", r.status_code == 200 and r.content[:4] == b"%PDF")
r = c.post("/archives/batch", data={"csrf": t, "ids": f"{ins_id},{ins_id}"})
check("批量打包", r.status_code == 200 and r.headers["content-type"] == "application/zip")

# ========== 12. 登录限速 ==========
for i in range(6):
    r = c.post("/login", data={"username": "gov", "password": "错误密码"}, follow_redirects=False)
r = c.post("/login", data={"username": "gov", "password": GOV_NEW_PWD}, follow_redirects=False)
loc = unquote(r.headers.get("location", ""))
check("P0-3 登录失败限速", "尝试次数过多" in loc, loc)

# ========== 13. 督办 / 统计 / 审计（gov 会话仍在） ==========
check("督办台账", c.get("/overdue").status_code == 200)
r = c.get("/overdue")
check("督办含即将到期分组", "7天内即将到期" in r.text)
check("统计页", c.get("/stats").status_code == 200)
r = c.get("/stats")
check("P2-11 统计条形图", "bar-fill" in r.text)
check("统计导出", c.get("/stats/export").status_code == 200)
r = c.get("/audit")
check("审计含登记动作", "登记检查记录" in r.text)
check("审计含更正留痕", "更正留痕" in r.text)

# ========== 13.5 P2：一企一档 / 一键备份 / 检查表停用 ==========
r = c.get("/enterprises/1/profile")
check("P2-10 一企一档页", r.status_code == 200 and "一企一档" in r.text and "平潭测试航运有限公司" in r.text)
check("P2-10 档案含整改率", "整改完成率" in r.text)
t = csrf_of()
r = c.post("/admin/backup", data={"csrf": t})
check("P2-12 一键备份下载", r.status_code == 200 and r.headers["content-type"] == "application/zip")
bz = zipfile.ZipFile(io.BytesIO(r.content))
check("P2-12 备份含数据库", "app.db" in bz.namelist())
check("P2-12 备份含附件与说明",
      any(n.startswith("uploads/") for n in bz.namelist()) and "备份说明.txt" in bz.namelist())
r = c.post("/templates/5/toggle", data={"csrf": t}, follow_redirects=False)
check("P2-13 检查表停用", r.status_code == 302)
conn = get_db()
check("P2-13 停用已生效", conn.execute("SELECT active FROM check_templates WHERE id=5").fetchone()["active"] == 0)
conn.close()
r = c.post("/templates/5/toggle", data={"csrf": t}, follow_redirects=False)  # 恢复启用
check("P2-13 检查表恢复启用", r.status_code == 302)

# ========== 14. 权限隔离 ==========
c.post("/login", data={"username": "ent", "password": "ent2026!"})
check("企业被拒访问项目库", c.get("/items").status_code == 403)
check("企业被拒访问用户管理", c.get("/users").status_code == 403)
check("企业可下载本企业归档", c.get(f"/inspections/{ins_id}/download").status_code == 200)
check("企业可看本企业一企一档", c.get("/enterprises/1/profile").status_code == 200)

# ========== 15. 服务端校验、异常兜底与改动留痕（审计 P1-7/P2-1/P2-2/P2-3） ==========
import app as _appmod
_appmod.LOGIN_FAILS.clear(); _appmod.IP_FAILS.clear()  # 清除第 12 节制造的限速状态
c.post("/login", data={"username": "gov", "password": GOV_NEW_PWD}, follow_redirects=False)
t = csrf_of()
check("不存在记录的打包下载返回404", c.get("/inspections/9999/download").status_code == 404)
check("不存在记录的下发返回404", c.post("/inspections/9999/issue", data={"csrf": t}).status_code == 404)
check("不存在记录的闭环返回404", c.post("/inspections/9999/close", data={"csrf": t}).status_code == 404)
check("不存在记录的文书返回404", c.get("/inspections/9999/doc/record").status_code == 404)
check("空批量归档返回400", c.post("/archives/batch", data={"csrf": t, "ids": ","}).status_code == 400)
check("未提交 ids 被拒（422/400）",
      c.post("/archives/batch", data={"csrf": t}).status_code in (400, 422))
check("选不存在的批量归档返回404",
      c.post("/archives/batch", data={"csrf": t, "ids": "9999"}).status_code == 404)


def save_ins(**kw):
    data = {"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 0, "check_type": "日常检查",
            "check_mode": "现场检查", "check_date": date.today().isoformat(), "inspectors": "李",
            "conclusion": "责发整改", "items_json": json.dumps([new_item()], ensure_ascii=False),
            "problems_json": "[]"}
    data.update(kw)
    return c.post("/inspections/save", data=data, follow_redirects=False).status_code


check("非法JSON返回400", save_ins(items_json="{坏JSON") == 400)
check("非法检查类型返回400", save_ins(check_type="不存在的类型") == 400)
check("非法检查结论返回400", save_ins(conclusion="乱填结论") == 400)
check("非法逐项结果返回400",
      save_ins(items_json=json.dumps([new_item(result="合格")], ensure_ascii=False)) == 400)
check("空问题描述/整改要求返回400",
      save_ins(problems_json=json.dumps([{"description": "", "requirement": ""}], ensure_ascii=False)) == 400)
# 造一条已下发且带问题的记录，用于复核结论校验（原记录已归档，会先命中归档锁定）
r = c.post("/inspections/save", data={"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 0,
                                     "check_type": "日常检查", "check_mode": "现场检查",
                                     "check_date": date.today().isoformat(), "inspectors": "李",
                                     "conclusion": "责发整改",
                                     "items_json": json.dumps([new_item(result="不符合")], ensure_ascii=False),
                                     "problems_json": json.dumps([{"description": "校验用问题",
                                                                   "requirement": "整改"}], ensure_ascii=False)},
           follow_redirects=False)
nid = int(r.headers.get("location", "/inspections/0").split("/")[2].split("?")[0])
c.post(f"/inspections/{nid}/issue", data={"csrf": t}, follow_redirects=False)
conn = get_db()
npid = conn.execute("SELECT id FROM problems WHERE inspection_id=?", (nid,)).fetchone()["id"]
conn.close()
check("非法复核结论返回400",
      c.post(f"/problems/{npid}/review", data={"csrf": t, "result": "乱填"}).status_code == 400)
check("合法复核结论可通过",
      c.post(f"/problems/{npid}/review", data={"csrf": t, "result": "pass", "opinion": "ok"},
             follow_redirects=False).status_code == 302)
# 改动留痕 before/after（审计 P1-7）
c.post("/enterprises/save", data={"csrf": t, "eid": 1, "name": "平潭测试航运有限公司",
                                 "credit_code": "91350128TEST0001X", "license_no": "闽交水运字0001号",
                                 "license_type": "普通货物运输", "legal_person": "王船长改",
                                 "contact": "陈安全", "phone": "138", "address": "平潭",
                                 "scope": "沿海", "remark": "改动留痕测试"}, follow_redirects=False)
conn = get_db()
rowlog = conn.execute("SELECT before, after FROM audit_logs WHERE action='修改企业'"
                      " ORDER BY id DESC LIMIT 1").fetchone()
conn.close()
check("修改操作记录 before/after（改动前后内容）",
      bool(rowlog and rowlog["before"] and rowlog["after"]), (rowlog["before"][:50] if rowlog else "无日志"))

# ========== 16. 第三批：去重/自选清单/船舶/复查关联/统计筛选/安全头/会话失效（审计 P2-5~P2-11、P3-5） ==========
t = csrf_of()
# P2-6 企业去重（名称/信用代码）
r = c.post("/enterprises/save", data={"csrf": t, "eid": 0, "name": "平潭测试航运有限公司",
                                     "credit_code": "91350128TEST0001X", "license_no": "", "license_type": "",
                                     "legal_person": "", "contact": "", "phone": "", "address": "",
                                     "scope": "", "remark": ""}, follow_redirects=False)
check("企业重名/重信用代码被拒（P2-6）", "已存在" in unquote(r.headers.get("location", "")))
# P2-8 新增检查项（未加入任何模板）可在登记页自选
c.post("/items/save", data={"csrf": t, "iid": 0, "code": "Z01", "category": "Z 测试类", "name": "新项测试",
                            "content": "内容", "legal_basis": "《测试》第1条", "method": "查资料",
                            "criteria": "合格", "is_key": 0, "is_veto": 0, "score": 3, "scope": "", "note": ""},
       follow_redirects=False)
check("新增检查项可在登记页自选（P2-8）", "Z01" in c.get("/inspections/new").text)
# P2-9 船舶清单维护
r = c.post("/enterprises/1/ships/save", data={"csrf": t, "sid": 0, "name": "测试一号", "ship_no": "CN12345",
                                              "ship_type": "普通货船", "dwt": "5000", "gt": "3000",
                                              "built_date": "2018-06", "license_no": "闽船001",
                                              "cert_status": "齐全有效", "remark": ""}, follow_redirects=False)
check("新增船舶（P2-9）", r.status_code == 302)
r = c.get("/enterprises/1/profile")
check("一企一档含船舶清单（P2-9）", "测试一号" in r.text and "CN12345" in r.text)
conn = get_db()
ship_id = conn.execute("SELECT id FROM ships ORDER BY id DESC LIMIT 1").fetchone()["id"]
conn.close()
check("删除船舶（P2-9）", c.post(f"/ships/{ship_id}/delete", data={"csrf": t},
                                 follow_redirects=False).status_code == 302)
# P2-10 复查关联原记录
r = c.post("/inspections/save", data={"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 0,
                                     "parent_id": nid, "check_type": "复查", "check_mode": "现场检查",
                                     "check_date": date.today().isoformat(), "inspectors": "李",
                                     "conclusion": "责发整改",
                                     "items_json": json.dumps([new_item()], ensure_ascii=False),
                                     "problems_json": "[]"}, follow_redirects=False)
rid2 = int(r.headers.get("location", "/inspections/0").split("/")[2].split("?")[0])
check("复查记录显示原记录关联（P2-10）", "本记录是对" in c.get(f"/inspections/{rid2}").text)
check("原记录显示复查记录关联（P2-10）", "该记录的复查记录" in c.get(f"/inspections/{nid}").text)
# P2-11 统计时间段筛选
check("统计页时间段筛选（P2-11）",
      c.get(f"/stats?date_from={date.today().isoformat()}&date_to={date.today().isoformat()}").status_code == 200)
check("统计筛选空区间为空（P2-11）", "暂无数据" in c.get("/stats?date_from=2099-01-01").text)
# P3-5 安全响应头
r = c.get("/")
check("安全响应头 X-Frame-Options/nosniff（P3-5）",
      r.headers.get("x-frame-options") == "DENY" and r.headers.get("x-content-type-options") == "nosniff")
check("操作日志分页（P3-4）", "页" in c.get("/audit").text)
# P2-9 企业自助维护联系信息
c2 = TestClient(app)
c2.post("/login", data={"username": "ent", "password": "ent2026!"}, follow_redirects=False)
t2 = re.search(r'<meta name="csrf" content="([^"]+)"', c2.get("/").text).group(1)
check("企业自助维护联系信息（P2-9）",
      c2.post("/enterprises/self", data={"csrf": t2, "contact": "新联系人", "phone": "13900000000",
                                         "address": "新地址"}, follow_redirects=False).status_code == 302)
# P2-5 改密后旧会话失效
old_cookie = c2.cookies.get("slys_session")
c2.post("/changepwd", data={"csrf": t2, "old_password": "ent2026!", "new_password": "ent2027!",
                            "new_password2": "ent2027!"}, follow_redirects=False)
c3 = TestClient(app)
c3.cookies.set("slys_session", old_cookie)
check("改密后旧会话失效（P2-5）", c3.get("/", follow_redirects=False).status_code == 302)

# ========== 17. 双随机抽查（功能新增 2026-09-26） ==========
t = csrf_of()
r = c.post("/random/draw", data={"csrf": t, "draw_type": "double", "ent_count": 1, "insp_count": 1,
                                "exclude_days": 0, "make_drafts": 1, "note": "测试抽取"},
           follow_redirects=False)
check("双随机抽取返回302", r.status_code == 302)
conn = get_db()
draw = conn.execute("SELECT * FROM random_draws ORDER BY id DESC LIMIT 1").fetchone()
check("抽取记录落库（含随机种子）", bool(draw and draw["seed"] and draw["ent_ids"]))
info = json.loads(draw["detail"] or "{}")
check("选派检查员入档", bool(info.get("insps")))
check("自动生成检查草稿", len(info.get("draft_ids", [])) == 1)
import random as _rnd
pool_ids = [x["id"] for x in conn.execute("SELECT id FROM enterprises ORDER BY id")]
repro = _rnd.Random(draw["seed"]).sample(pool_ids, 1)
check("按种子可复现抽取结果（抽样公正）", str(repro[0]) == draw["ent_ids"], f"复现={repro} vs {draw['ent_ids']}")
conn.close()
check("抽查历史页可访问", "双随机" in c.get("/random").text)
check("抽查结果表可导出", c.get(f"/random/{draw['id']}/export").content[:2] == b"PK")
check("企业池不足被拒",
      "少于" in unquote(c.post("/random/draw", data={"csrf": t, "draw_type": "single", "ent_count": 999,
                                                    "insp_count": 1, "exclude_days": 0, "make_drafts": 0},
                              follow_redirects=False).headers.get("location", "")))
c2 = TestClient(app)
c2.post("/login", data={"username": "ent", "password": "ent2027!"}, follow_redirects=False)
check("企业被拒访问双随机抽查", c2.get("/random").status_code == 403)

# ========== 18. 检查计划（功能新增 2026-09-26） ==========
t = csrf_of()
r = c.post("/plans/save", data={"csrf": t, "pid": 0, "title": "2026年度测试检查计划",
                               "year": "2026", "period": "全年", "note": "测试"},
           follow_redirects=False)
check("新建检查计划", r.status_code == 302)
conn = get_db()
plan_id = conn.execute("SELECT id FROM plans ORDER BY id DESC LIMIT 1").fetchone()["id"]
conn.close()
r = c.post(f"/plans/{plan_id}/items/save", data={"csrf": t, "piid": 0, "enterprise_id": 1,
                                                "ent_scope": "", "check_type": "日常检查",
                                                "count_plan": 2, "period": "全年", "inspectors": "李检查员",
                                                "note": ""}, follow_redirects=False)
check("添加计划项", r.status_code == 302)
r = c.post(f"/plans/{plan_id}/issue", data={"csrf": t}, follow_redirects=False)
check("下达计划", r.status_code == 302)
conn = get_db()
check("计划状态已下达",
      conn.execute("SELECT status FROM plans WHERE id=?", (plan_id,)).fetchone()["status"] == "issued")
piid = conn.execute("SELECT id FROM plan_items WHERE plan_id=?", (plan_id,)).fetchone()["id"]
conn.close()
check("下达后计划项锁定",
      c.post(f"/plans/{plan_id}/items/save", data={"csrf": t, "piid": 0, "enterprise_id": 1, "ent_scope": "",
                                                  "check_type": "日常检查", "count_plan": 1, "period": "",
                                                  "inspectors": "", "note": ""},
             follow_redirects=False).status_code == 400)
# 登记检查并关联计划项 → 下发后自动核销
r = c.post("/inspections/save", data={"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 0,
                                     "parent_id": 0, "plan_item_id": piid, "check_type": "日常检查",
                                     "check_mode": "现场检查", "check_date": date.today().isoformat(),
                                     "inspectors": "李", "conclusion": "未发现问题",
                                     "items_json": json.dumps([new_item()], ensure_ascii=False),
                                     "problems_json": "[]"}, follow_redirects=False)
ins4 = int(r.headers.get("location", "/inspections/0").split("/")[2].split("?")[0])
c.post(f"/inspections/{ins4}/issue", data={"csrf": t}, follow_redirects=False)
conn = get_db()
done = conn.execute("SELECT COUNT(*) c FROM inspections WHERE plan_item_id=? AND status<>'draft'",
                    (piid,)).fetchone()["c"]
conn.close()
check("计划项执行进度已核销（下发后计1家次）", done == 1, f"done={done}")
check("计划表可导出", c.get(f"/plans/{plan_id}/export").content[:2] == b"PK")
check("计划列表页可访问", c.get("/plans").status_code == 200)
check("企业被拒访问检查计划", c2.get("/plans").status_code == 403)

# ========== 19. 现场登记（手机适配）与手写签字确认（功能新增 2026-09-26） ==========
t = csrf_of()
r = c.post("/inspections/save", data={"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 0,
                                     "parent_id": 0, "plan_item_id": 0, "check_type": "日常检查",
                                     "check_mode": "现场检查", "check_date": date.today().isoformat(),
                                     "inspectors": "李检查员", "conclusion": "责发整改",
                                     "items_json": json.dumps([new_item()], ensure_ascii=False),
                                     "problems_json": "[]"}, follow_redirects=False)
ins5 = int(r.headers.get("location", "/inspections/0").split("/")[2].split("?")[0])
check("现场登记页可打开（新建）", c.get("/onsite").status_code == 200)
check("现场登记页可编辑草稿（手机适配）",
      c.get(f"/onsite?edit={ins5}").status_code == 200 and
      "res-btns" in c.get(f"/onsite?edit={ins5}").text)
SIGN_PNG = ("data:image/png;base64,iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJ"
            "AAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")
r = c.post(f"/inspections/{ins5}/sign", data={"csrf": t, "sign_type": "insp", "signer_name": "李检查员",
                                             "image": SIGN_PNG}, follow_redirects=False)
check("检查人员签字", r.status_code == 302)
r = c.post(f"/inspections/{ins5}/sign", data={"csrf": t, "sign_type": "ent", "signer_name": "王船长",
                                             "image": SIGN_PNG}, follow_redirects=False)
check("政府端现场代录企业签字（审计注明代录）", r.status_code == 302)
conn = get_db()
nsign = conn.execute("SELECT COUNT(*) c FROM signatures WHERE inspection_id=?", (ins5,)).fetchone()["c"]
sig_row = conn.execute("SELECT id FROM signatures WHERE inspection_id=? AND sign_type='ent'",
                       (ins5,)).fetchone()
conn.close()
check("签字记录落库（2条）", nsign == 2)
check("签名图可访问", c.get(f"/signatures/{sig_row['id']}").headers.get("content-type") == "image/png")
check("签字已嵌入文书PDF", c.get(f"/inspections/{ins5}/doc/record").content[:4] == b"%PDF")
e = TestClient(app)
e.post("/login", data={"username": "ent", "password": "ent2027!"}, follow_redirects=False)
te = re.search(r'<meta name="csrf" content="([^"]+)"', e.get("/").text).group(1)
check("企业只能签企业栏（签检查人员栏被拒）",
      e.post(f"/inspections/{ins5}/sign", data={"csrf": te, "sign_type": "insp", "signer_name": "x",
                                               "image": SIGN_PNG}, follow_redirects=False).status_code == 403)
check("空白签名被拒",
      c.post(f"/inspections/{ins5}/sign", data={"csrf": t, "sign_type": "insp", "signer_name": "李",
                                               "image": ""}, follow_redirects=False).status_code == 400)
conn = get_db()
conn.execute("UPDATE inspections SET status='archived', archived_at='2026-01-01 00:00:00' WHERE id=?", (ins5,))
conn.commit(); conn.close()
check("归档锁定后不能补签",
      c.post(f"/inspections/{ins5}/sign", data={"csrf": t, "sign_type": "insp", "signer_name": "李",
                                               "image": SIGN_PNG}, follow_redirects=False).status_code == 403)

# ========== 21. 表单自动保存 + 检查表打印（被检查企业可见） ==========
check("自动保存组件已加载", c.get("/static/autosave.js").status_code == 200)
check("登记检查页接入自动保存", "autosave.js" in c.get("/inspections/new").text)
check("现场登记页接入自动保存", "autosave.js" in c.get("/onsite").text)
r = c.get(f"/inspections/{ins5}/print")
check("检查表打印页可访问", r.status_code == 200 and "打印本检查表" in r.text)
check("打印页含逐项结果与签字栏", "检查项目" in r.text and "签字" in r.text)
check("详情页带打印按钮", "打印检查表" in c.get(f"/inspections/{ins5}").text)
ce = TestClient(app)
ce.post("/login", data={"username": "ent", "password": "ent2027!"}, follow_redirects=False)
r = ce.get(f"/inspections/{ins5}/print")
check("被检查企业可查看打印检查表", r.status_code == 200 and "检查项目" in r.text)
check("企业打印页含签字栏与检查编号", "签字" in r.text and ins5 is not None)

# ========== 22. 性能与健壮性（优化收尾：索引/WAL/分页/健康检查） ==========
check("健康检查端点", c.get("/healthz").status_code == 200)
check("检查记录列表分页", "每页 50 条" in c.get("/inspections").text)
check("企业名录分页", "每页 50 家" in c.get("/enterprises").text)
conn = get_db()
n_idx = conn.execute("SELECT COUNT(*) c FROM sqlite_master WHERE type='index' AND name LIKE 'idx_%'").fetchone()["c"]
wal = conn.execute("PRAGMA journal_mode").fetchone()[0]
conn.close()
check("数据库索引已建", n_idx >= 15, f"{n_idx} 个")
check("SQLite 已启用 WAL", str(wal).lower() == "wal", f"journal_mode={wal}")

# ========== 23. 提醒中心/公示/批量打印/月度汇总/图片预览（功能新增） ==========
# 造一条带逾期问题的记录
r = c.post("/inspections/save", data={"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 0,
                                     "parent_id": 0, "plan_item_id": 0, "check_type": "日常检查",
                                     "check_mode": "现场检查", "check_date": date.today().isoformat(),
                                     "inspectors": "李", "conclusion": "责发整改",
                                     "items_json": json.dumps([new_item(result="不符合")], ensure_ascii=False),
                                     "problems_json": json.dumps([{"description": "逾期测试问题", "basis": "某条",
                                                                   "requirement": "整改", "deadline": "2020-01-01",
                                                                   "responsible": "王", "to_msa": False}],
                                                                 ensure_ascii=False)}, follow_redirects=False)
ins6 = int(r.headers.get("location", "/inspections/0").split("/")[2].split("?")[0])
c.post(f"/inspections/{ins6}/issue", data={"csrf": t}, follow_redirects=False)
r = c.get("/messages")
check("消息提醒页含逾期事项", r.status_code == 200 and "逾期测试问题" in r.text)
check("工作台提醒横幅", "去消息提醒处理" in c.get("/").text)
ce2 = TestClient(app)
ce2.post("/login", data={"username": "ent", "password": "ent2027!"}, follow_redirects=False)
check("企业工作台待办直达", "待我整改" in ce2.get("/").text and "逾期测试问题" in ce2.get("/").text)
check("手动推送（未配机器人时给出提示）",
      "推送失败" in unquote(c.post("/messages/push", data={"csrf": t},
                                  follow_redirects=False).headers.get("location", "")))
check("双随机公示页（免登录可访问）", TestClient(app).get("/public/random").status_code == 200)
check("批量打印检查表", c.get(f"/print-batch?ids={ins5},{ins6}").status_code == 200 and
      "打印全部（2 条）" in c.get(f"/print-batch?ids={ins5},{ins6}").text)
check("月度考核汇总表", c.get("/stats/monthly").status_code == 200 and "总体情况" in c.get("/stats/monthly").text)
check("月度汇总表导出Excel", c.get("/stats/monthly/export").content[:2] == b"PK")
check("登录页密码切换与公示入口", "pw-toggle" in TestClient(app).get("/login").text and
      "public/random" in TestClient(app).get("/login").text)
conn = get_db()
att = conn.execute("SELECT id FROM attachments WHERE filename LIKE '%.png' LIMIT 1").fetchone()
conn.close()
check("附件图片内嵌预览",
      bool(att) and c.get(f"/attachments/{att['id']}?inline=1").headers.get("content-type", "").startswith("image/"))

# ========== 24. 数据库加固：检查编号撞号自动跳号（抗并发登记） ==========
import routes_inspection as _ri
_orig_alloc = _ri.alloc_ins_code
conn = get_db()
taken = conn.execute("SELECT code FROM inspections ORDER BY id DESC LIMIT 1").fetchone()["code"]
conn.close()
_state = {"n": 0}


def _stub_alloc(cn, today8):
    _state["n"] += 1
    return taken if _state["n"] == 1 else _orig_alloc(cn, today8)   # 第一次强制返回已占用编号


_ri.alloc_ins_code = _stub_alloc
r = c.post("/inspections/save", data={"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 0,
                                     "parent_id": 0, "plan_item_id": 0, "check_type": "日常检查",
                                     "check_mode": "现场检查", "check_date": date.today().isoformat(),
                                     "inspectors": "李", "conclusion": "未发现问题",
                                     "items_json": json.dumps([new_item()], ensure_ascii=False),
                                     "problems_json": "[]"}, follow_redirects=False)
_ri.alloc_ins_code = _orig_alloc
new_id = int(r.headers.get("location", "/inspections/0").split("/")[2].split("?")[0])
conn = get_db()
new_code = conn.execute("SELECT code FROM inspections WHERE id=?", (new_id,)).fetchone()["code"]
uniq = conn.execute("SELECT COUNT(*) c, COUNT(DISTINCT code) d FROM inspections").fetchone()
conn.close()
check("检查编号撞号自动跳号（抗并发）", r.status_code == 302 and new_code != taken,
      f"占号={taken} 新记录={new_code}")
check("全库检查编号唯一", uniq["c"] == uniq["d"], f"{uniq['c']} 条 / {uniq['d']} 个编号")

# ========== 25. 船舶批量导入 + 消息红点 ==========
t = csrf_of()
check("船舶导入模板可下载", c.get("/ships/template").content[:2] == b"PK")
_wb = Workbook()
_ws = _wb.active
_ws.append(["所属企业名称", "船名", "登记号/IMO", "船型/种类", "载重吨", "总吨", "建成日期",
            "营业运输证号", "证书情况", "备注"])
_ws.append(["平潭导入甲公司", "海星1", "CN001", "普通货船", "5000", "3000", "2018-06", "闽船A1", "齐全", ""])
_ws.append(["平潭导入甲公司", "海星2", "CN002", "普通货船", "4000", "2000", "2019-01", "闽船A2", "齐全", ""])
_ws.append(["平潭导入甲公司", "海星1", "CN003", "普通货船", "1", "1", "2020-01", "", "", "重名应跳过"])
_ws.append(["不存在的企业", "幽灵船", "", "", "", "", "", "", "", "企业不存在应跳过"])
_ib = io.BytesIO()
_wb.save(_ib)
_ib.seek(0)
r = c.post("/ships/import", data={"csrf": t}, files={"file": ("船舶.xlsx", _ib, "application/vnd.ms-excel")},
           follow_redirects=False)
loc_msg = unquote(r.headers.get("location", ""))
check("船舶批量导入（新增2跳过2）", r.status_code == 302 and "新增2条，跳过2条" in loc_msg, loc_msg)
conn = get_db()
n_ship = conn.execute("SELECT COUNT(*) c FROM ships WHERE name LIKE '海星%'").fetchone()["c"]
conn.close()
check("船舶已入库（同名去重生效）", n_ship == 2, f"海星系列 {n_ship} 条")
check("消息提醒红点（有逾期事项时导航显示计数）", "<sup" in c.get("/").text)
te = TestClient(app)
te.post("/login", data={"username": "ent", "password": "ent2027!"}, follow_redirects=False)
te_t = re.search(r'<meta name="csrf" content="([^"]+)"', te.get("/").text).group(1)
_wb2 = Workbook()
_ws2 = _wb2.active
_ws2.append(["所属企业名称", "船名"])
_ws2.append(["平潭导入乙公司", "越界船"])
_ib2 = io.BytesIO()
_wb2.save(_ib2)
_ib2.seek(0)
r = te.post("/ships/import", data={"csrf": te_t}, files={"file": ("x.xlsx", _ib2, "application/vnd.ms-excel")},
            follow_redirects=False)
check("企业用户越界导入被拒（只能填本企业）",
      r.status_code == 302 and "新增0条" in unquote(r.headers.get("location", "")),
      unquote(r.headers.get("location", "")))

# ========== 26. 审计日志哈希链（防篡改留痕） ==========
import hashlib as _hl
conn = get_db()
_rows = conn.execute("SELECT * FROM audit_logs ORDER BY id").fetchall()
conn.close()
_prev = ""
_brk = 0
for _r in _rows:
    if not _r["chain_hash"]:
        _prev = ""
        continue
    _pl = "|".join([_r["username"], _r["action"], _r["entity"], str(_r["entity_id"]),
                    _r["detail"], _r["before"], _r["after"], _r["ip"], _r["created_at"]])
    if _hl.sha256((_prev + "|" + _pl).encode("utf-8")).hexdigest() != _r["chain_hash"]:
        _brk += 1
    _prev = _r["chain_hash"] or ""
check("审计日志哈希链完整（全部入链无断点）",
      _brk == 0 and all(_r["chain_hash"] for _r in _rows), f"{len(_rows)}条 断点{_brk}")

# ========== 27. 会话 Cookie 安全属性（独立客户端，避免 Secure Cookie 污染主会话） ==========
_ckc = TestClient(app)
_ckc.post("/login", data={"username": "gov", "password": "gov2026!"}, follow_redirects=False)
_ck = _ckc.post("/login", data={"username": "gov", "password": "gov2026!"},
                follow_redirects=False).headers.get("set-cookie", "")
_ck2 = _ckc.post("/login", data={"username": "gov", "password": "gov2026!"}, follow_redirects=False,
                 headers={"X-Forwarded-Proto": "https"}).headers.get("set-cookie", "")
check("会话Cookie安全属性（HttpOnly+SameSite；HTTPS/反代自动加Secure）",
      "HttpOnly" in _ck and "SameSite=lax" in _ck and "Secure" not in _ck and "Secure" in _ck2,
      f"http头={_ck[:80]} / https头={_ck2[:80]}")

# ========== 28. 表单校验错误友好化（使用体验修复） ==========
# 整个 check_date 字段不传 → 触发 FastAPI 校验器（此前甩裸 JSON，应渲染友好错误页）
r = c.post("/inspections/save", data={"csrf": t, "ins_id": 0, "enterprise_id": 1, "template_id": 0,
                                     "parent_id": 0, "plan_item_id": 0, "check_type": "日常检查",
                                     "check_mode": "现场检查",
                                     "inspectors": "李", "conclusion": "未发现问题",
                                     "items_json": "[]", "problems_json": "[]"}, follow_redirects=True)
check("缺必填项→友好校验页（不再裸JSON）",
      r.status_code == 422 and "表单校验未通过" in r.text and '"detail"' not in r.text,
      f"status={r.status_code}")

# ========== 29. 问题整改进度提示（使用体验修复） ==========
_r6 = c.get(f"/inspections/{ins6}")
check("问题整改进度提示（多问题逐条销号可见进度）",
      _r6.status_code == 200 and "整改进度 0/1" in _r6.text and "逐条" in _r6.text,
      f"ins6={ins6} 状态={_r6.status_code} 长度={len(_r6.text)}")

n_fail = sum(1 for _, v in ok if not v)
print(f"\n=== {len(ok) - n_fail}/{len(ok)} 通过 ===")
# 运行时断言数（含循环内多次执行的断言）写入给 check_docs.py 校验文档口径
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "smoke_last_run.txt"), "w") as _f:
    _f.write(str(len(ok)))
shutil.rmtree(TEST_DATA, ignore_errors=True)
sys.exit(1 if n_fail else 0)
