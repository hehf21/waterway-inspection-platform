# -*- coding: utf-8 -*-
"""归档打包：单条/批量 ZIP，内含 PDF 文书、附件、xlsx 目录、JSON、SHA256SUMS"""
import hashlib
import io
import json
import os
import re
import sqlite3
import zipfile
from datetime import datetime

from openpyxl import Workbook

import pdfgen
from config import ARCHIVE_DIR, UPLOAD_DIR
from db import audit

SAFE_CHARS = re.compile(r'[\\/:*?"<>|\s]+')
SAFE_FILENAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]+')


def safe_name(name: str) -> str:
    return SAFE_CHARS.sub("_", str(name)).strip("_")[:40] or "未命名"


def safe_filename(name: str, limit: int = 120) -> str:
    """归档包落盘/下载用文件名：剔除路径分隔符与控制字符，避免 ZIP 条目路径穿越。"""
    return SAFE_FILENAME.sub("_", str(name)).strip(" ._")[:limit] or "归档包"


def _pack_path(file_name: str) -> str:
    """归档包在 data\\archives 下的绝对路径。"""
    return os.path.join(ARCHIVE_DIR, os.path.basename(file_name))


def _load_existing_pack(conn, inspection_id: int):
    """若该记录的归档包已落盘则直接复用，保证同一归档件字节一致、校验值可复核。"""
    row = conn.execute("SELECT archive_file FROM inspections WHERE id=?", (inspection_id,)).fetchone()
    if not row or not row["archive_file"]:
        return None
    path = _pack_path(row["archive_file"])
    if not os.path.exists(path):
        return None
    with open(path, "rb") as fp:
        data = fp.read()
    cnt = conn.execute("SELECT file_count FROM archive_logs WHERE file_name=? ORDER BY id DESC LIMIT 1",
                       (row["archive_file"],)).fetchone()
    return data, row["archive_file"], _sha256(data), len(data), (cnt["file_count"] if cnt else 0)


def _sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def _ent_short(name: str) -> str:
    n = re.sub(r"(有限公司|股份有限公司|有限责任公司|公司|平潭|福建省|福建)", "", name)
    return n[:8] or name[:8]


def build_record_pack(conn: sqlite3.Connection, inspection_id: int):
    """构建单条记录归档包，返回 (zip_bytes, zip_name, sha256, size, file_count)"""
    ins = conn.execute("SELECT * FROM inspections WHERE id=?", (inspection_id,)).fetchone()
    if not ins:
        raise ValueError("检查记录不存在")
    ent = conn.execute("SELECT * FROM enterprises WHERE id=?", (ins["enterprise_id"],)).fetchone()
    items = conn.execute("SELECT * FROM inspection_items WHERE inspection_id=? ORDER BY sort_order",
                         (inspection_id,)).fetchall()
    problems = conn.execute("SELECT * FROM problems WHERE inspection_id=? ORDER BY seq", (inspection_id,)).fetchall()

    from config import REC_STATUS_CN
    info = {
        "code": ins["code"], "ent_name": ent["name"] if ent else "",
        "credit_code": ent["credit_code"] if ent else "", "license_no": ent["license_no"] if ent else "",
        "check_date": ins["check_date"], "check_type": ins["check_type"], "check_mode": ins["check_mode"],
        "location": ins["location"], "inspectors": ins["inspectors"], "conclusion": ins["conclusion"],
        "conclusion_note": ins["conclusion_note"], "deadline": ins["deadline"],
        "status_cn": REC_STATUS_CN.get(ins["status"], ins["status"]),
    }

    def att(owner_type, owner_id):
        return conn.execute(
            "SELECT * FROM attachments WHERE owner_type=? AND owner_id=? ORDER BY id", (owner_type, owner_id)).fetchall()

    files = []  # (arcname, bytes)

    # 手写签名图（被检查企业 / 检查人员），嵌入文书签字栏
    signs = {}
    for s in conn.execute("SELECT * FROM signatures WHERE inspection_id=?", (inspection_id,)).fetchall():
        p = os.path.join(UPLOAD_DIR, s["stored_name"])
        if os.path.exists(p):
            signs[s["sign_type"]] = p

    # 01 现场检查记录表
    buf = io.BytesIO()
    pdfgen.record_pdf(info, [dict(r) for r in items], buf, signs=signs)
    files.append(("01-现场检查记录表.pdf", buf.getvalue()))

    # 02 问题清单及整改要求
    prob_rows = []
    for p in problems:
        d = dict(p)
        d["seq_label"] = f"P{p['seq']}"
        prob_rows.append(d)
    buf = io.BytesIO()
    pdfgen.notice_pdf(info, prob_rows, buf, signs=signs)
    files.append(("02-问题清单及整改要求.pdf", buf.getvalue()))

    # 03 企业整改情况报告
    fb_rows = []
    for p in prob_rows:
        fbs = conn.execute("SELECT * FROM feedbacks WHERE problem_id=? ORDER BY round DESC", (p["id"],)).fetchall()
        f = fbs[0] if fbs else None
        ats = att("problem", p["id"]) + (att("feedback", f["id"]) if f else [])
        fb_rows.append({
            "seq_label": p["seq_label"], "description": p["description"],
            "reason": f["reason"] if f else "（未反馈）",
            "measure": f["measure"] if f else "（未反馈）",
            "completion": f["completion"] if f else "", "done_date": f["done_date"] if f else "",
            "submitter": f["submitter"] if f else "",
            "att_names": "、".join(a["filename"] for a in ats) or "无",
        })
    buf = io.BytesIO()
    pdfgen.feedback_pdf(info, fb_rows, buf, signs=signs)
    files.append(("03-企业整改情况报告.pdf", buf.getvalue()))

    # 04 附件（原始附件按问题分子文件夹）
    att_count = 0
    for p in prob_rows:
        ats = att("problem", p["id"])
        fbs = conn.execute("SELECT id FROM feedbacks WHERE problem_id=? ORDER BY round DESC", (p["id"],)).fetchall()
        for f in fbs:
            ats += att("feedback", f["id"])
        for a in ats:
            src = os.path.join(UPLOAD_DIR, a["stored_name"])
            if os.path.exists(src):
                with open(src, "rb") as fp:
                    files.append((f"04-附件/{p['seq_label']}-{safe_name(p['description'])[:12]}/{a['filename']}", fp.read()))
                att_count += 1
    ins_ats = att("inspection", inspection_id)
    for a in ins_ats:
        src = os.path.join(UPLOAD_DIR, a["stored_name"])
        if os.path.exists(src):
            with open(src, "rb") as fp:
                files.append((f"04-附件/00-检查现场材料/{a['filename']}", fp.read()))
            att_count += 1

    # 05 复核意见与复查记录
    rv_rows = []
    for p in prob_rows:
        rvs = conn.execute("SELECT * FROM reviews WHERE problem_id=? ORDER BY reviewed_at DESC", (p["id"],)).fetchall()
        for r in rvs:
            rv_rows.append({
                "seq_label": p["seq_label"], "summary": (r["opinion"] or "")[:60],
                "result_cn": "通过" if r["result"] == "pass" else "退回",
                "opinion": r["opinion"], "recheck_date": r["recheck_date"],
                "recheck_note": r["recheck_note"], "reviewer": r["reviewer"],
            })
    buf = io.BytesIO()
    pdfgen.review_pdf(info, rv_rows, buf)
    files.append(("05-复核意见与复查记录.pdf", buf.getvalue()))

    # 07 结构化数据（先生成，目录要用到）
    record_json = {
        "inspection": {**info, "id": inspection_id, "created_by": ins["created_by"],
                       "issued_at": ins["issued_at"], "closed_at": ins["closed_at"]},
        "items": [dict(r) for r in items],
        "problems": prob_rows,
        "feedbacks": [{"problem_seq": p["seq"], **dict(f)}
                      for p in problems
                      for f in conn.execute("SELECT * FROM feedbacks WHERE problem_id=? ORDER BY round", (p["id"],))],
        "reviews": [{"problem_seq": p["seq"], **dict(r)}
                    for p in problems
                    for r in conn.execute("SELECT * FROM reviews WHERE problem_id=? ORDER BY reviewed_at", (p["id"],))],
        "attachments": [{"problem_seq": p["seq"], **dict(a)}
                        for p in problems for a in att("problem", p["id"])]
                       + [{"problem_seq": None, **dict(a)} for a in ins_ats],
    }
    json_bytes = json.dumps(record_json, ensure_ascii=False, indent=2).encode("utf-8")
    files.append(("07-结构化数据.json", json_bytes))

    wb = Workbook()
    ws = wb.active
    ws.title = "检查记录"
    ws.append(["检查编号", "企业名称", "检查日期", "检查类型", "检查方式", "检查人员", "检查结论", "状态"])
    ws.append([info["code"], info["ent_name"], info["check_date"], info["check_type"], info["check_mode"],
               info["inspectors"], info["conclusion"], info["status_cn"]])
    ws2 = wb.create_sheet("检查项结果")
    ws2.append(["序号", "检查项编号", "类别", "检查项目", "检查内容", "检查依据", "检查方法", "检查结果", "情况记录"])
    for r in items:
        ws2.append([r["sort_order"], r["item_code"], r["category"], r["item_name"], r["item_content"],
                    r["legal_basis"], r["method"], r["result"], r["finding"]])
    ws3 = wb.create_sheet("问题与整改")
    ws3.append(["问题编号", "问题描述", "违反条款", "整改要求", "整改期限", "责任人", "状态", "整改措施", "完成情况", "复核结论", "复核意见"])
    for p in prob_rows:
        fbs = conn.execute("SELECT * FROM feedbacks WHERE problem_id=? ORDER BY round DESC", (p["id"],)).fetchall()
        rvs = conn.execute("SELECT * FROM reviews WHERE problem_id=? ORDER BY reviewed_at DESC", (p["id"],)).fetchall()
        f, r = (fbs[0] if fbs else None), (rvs[0] if rvs else None)
        ws3.append([p["seq_label"], p["description"], p["legal_basis"], p["requirement"], p["deadline"],
                    p["responsible"], p["status"], f["measure"] if f else "", f["completion"] if f else "",
                    r["result"] if r else "", r["opinion"] if r else ""])
    xls_buf = io.BytesIO()
    wb.save(xls_buf)
    files.append(("07-结构化数据.xlsx", xls_buf.getvalue()))

    # 06 归档目录（xlsx + 说明）
    catalog = [
        {"no": 1, "name": "01-现场检查记录表.pdf", "kind": "检查文书", "size": f"{len(files[0][1])} 字节", "note": "含逐项检查结果、双方签字栏"},
        {"no": 2, "name": "02-问题清单及整改要求.pdf", "kind": "检查文书", "size": f"{len(files[1][1])} 字节", "note": "问题、违反条款、整改要求、期限"},
        {"no": 3, "name": "03-企业整改情况报告.pdf", "kind": "整改材料", "size": f"{len(files[2][1])} 字节", "note": "企业逐项整改反馈"},
        {"no": 4, "name": "04-附件/", "kind": "原始附件", "size": f"{att_count} 件", "note": "按问题分子文件夹，保留原文件名"},
        {"no": 5, "name": "05-复核意见与复查记录.pdf", "kind": "复核文书", "size": f"{len(rv_rows)} 条", "note": "复核结论、意见、复查情况"},
        {"no": 6, "name": "06-归档目录.xlsx / 归档说明.txt", "kind": "目录", "size": "", "note": "文件清单与归档说明"},
        {"no": 7, "name": "07-结构化数据.json / .xlsx", "kind": "电子数据", "size": "", "note": "便于系统迁移与再利用"},
        {"no": 8, "name": "SHA256SUMS.txt", "kind": "校验", "size": "", "note": "全包文件完整性校验值"},
    ]
    wb2 = Workbook()
    ws = wb2.active
    ws.title = "归档目录"
    ws.append(["序号", "文件名", "类别", "件数/大小", "备注"])
    for c in catalog:
        ws.append([c["no"], c["name"], c["kind"], c["size"], c["note"]])
    ws.append([])
    ws.append(["归档单位", "平潭综合实验区交通与建设局", "", "", ""])
    ws.append(["归档时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S"), "", "", ""])
    cat_buf = io.BytesIO()
    wb2.save(cat_buf)
    files.append(("06-归档目录.xlsx", cat_buf.getvalue()))

    readme = (f"检查记录归档说明\n{'='*40}\n"
              f"检查编号：{info['code']}\n被检查企业：{info['ent_name']}\n"
              f"检查日期：{info['check_date']}　检查类型：{info['check_type']}\n"
              f"检查人员：{info['inspectors']}\n检查结论：{info['conclusion']}\n"
              f"记录状态：{info['status_cn']}\n"
              f"归档时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
              f"生成系统：水路运输企业检查记录管理平台\n"
              f"完整性校验：请核对 SHA256SUMS.txt 中各文件 SHA-256 校验值。\n")
    files.append(("06-归档说明.txt", readme.encode("utf-8")))

    # SHA256SUMS
    sums = "".join(f"{_sha256(data)}  {name}\n" for name, data in files)
    files.append(("SHA256SUMS.txt", sums.encode("utf-8")))

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files:
            z.writestr(name, data)
    zip_bytes = buf.getvalue()

    ent_name = _ent_short(ent["name"]) if ent else "企业"
    date_part = ins["check_date"].replace("-", "")[:8] if ins["check_date"] else ""
    zip_name = safe_filename(f"{ent_name}_{date_part}_{ins['check_type']}_{ins['code']}.zip")
    return zip_bytes, zip_name, _sha256(zip_bytes), len(zip_bytes), len(files)


def archive_and_download(conn, inspection_ids, username, ip="", archive=False):
    """单条或批量打包。返回 (zip_bytes, zip_name, sha256, size, file_count)

    - archive=True 才把已闭环记录转“已归档”并锁定，仅由显式归档动作触发；默认只读、不改状态（审计 P1-2）；
    - 归档包一律落盘到 ARCHIVE_DIR；单条记录再次下载复用同一份包，保证字节一致、校验值可复核（审计 P1-1）。
    """
    if not inspection_ids:
        raise ValueError("未选择检查记录")
    scope = "single" if len(inspection_ids) == 1 else "batch"
    reused = False
    if scope == "single":
        existing = _load_existing_pack(conn, inspection_ids[0])
        if existing:
            zip_bytes, zip_name, sha, size, cnt = existing
            reused = True

    if not reused:
        if scope == "single":
            zip_bytes, zip_name, sha, size, cnt = build_record_pack(conn, inspection_ids[0])
        else:
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
                # 总目录
                wb = Workbook()
                ws = wb.active
                ws.title = "总目录"
                ws.append(["序号", "检查编号", "企业名称", "检查日期", "检查类型", "结论", "状态", "归档子文件夹"])
                for i, rid in enumerate(inspection_ids, 1):
                    b, n, _s, _sz, _c = build_record_pack(conn, rid)
                    ins = conn.execute("SELECT * FROM inspections WHERE id=?", (rid,)).fetchone()
                    if not ins:
                        continue
                    ent = conn.execute("SELECT name FROM enterprises WHERE id=?", (ins["enterprise_id"],)).fetchone()
                    sub = f"{i:02d}-{ins['code']}"
                    z.writestr(f"{sub}/{n}", b)
                    ws.append([i, ins["code"], ent["name"] if ent else "", ins["check_date"],
                               ins["check_type"], ins["conclusion"], ins["status"], f"{sub}/"])
                tbuf = io.BytesIO()
                wb.save(tbuf)
                z.writestr("总目录.xlsx", tbuf.getvalue())
                z.writestr("批量归档说明.txt",
                           f"批量归档包：共 {len(inspection_ids)} 条检查记录\n生成时间："
                           f"{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n生成人：{username}\n"
                           f"各子文件夹内含该条记录完整归档材料与 SHA256SUMS.txt。\n".encode("utf-8"))
            zip_bytes = buf.getvalue()
            zip_name = safe_filename(f"水路运输检查记录批量归档_{datetime.now().strftime('%Y%m%d_%H%M%S')}.zip")
            sha, size = _sha256(zip_bytes), len(zip_bytes)
            # 件数按包内实际条目数计（子包 + 总目录 + 说明），审计 P2-7
            cnt = len(zipfile.ZipFile(io.BytesIO(zip_bytes)).namelist())

        # 归档包落盘，保证“已归档”记录随时可复现同一份成品（审计 P1-1）
        os.makedirs(ARCHIVE_DIR, exist_ok=True)
        with open(_pack_path(zip_name), "wb") as fp:
            fp.write(zip_bytes)
        conn.execute(
            "INSERT INTO archive_logs(scope,record_ids,file_name,sha256,size,file_count,generated_by,"
            "file_path,generated_at) VALUES(?,?,?,?,?,?,?,?,?)",
            (scope, ",".join(str(i) for i in inspection_ids), zip_name, sha, size, cnt, username,
             _pack_path(zip_name), datetime.now().strftime("%Y-%m-%d %H:%M:%S")))

    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    for rid in inspection_ids:
        row = conn.execute("SELECT status FROM inspections WHERE id=?", (rid,)).fetchone()
        if not row:
            continue
        if archive and row["status"] == "closed":
            conn.execute("UPDATE inspections SET status='archived', archived_at=?, archive_file=? WHERE id=?",
                         (now, zip_name, rid))
        else:
            conn.execute("UPDATE inspections SET archive_file=? WHERE id=?", (zip_name, rid))
    audit(conn, username, "复用已归档包下载" if reused else "打包归档下载", "archive",
          ",".join(map(str, inspection_ids)),
          f"{'复用' if reused else '生成'}归档包 {zip_name}（{cnt} 件，落盘 {_pack_path(zip_name)}）", ip=ip)
    conn.commit()
    return zip_bytes, zip_name, sha, size, cnt
