# -*- coding: utf-8 -*-
"""检查项目库 / 检查表导出（Excel / Word），供逐条核对法规条款"""
import io

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook

HEAD = ["编号", "类别", "检查项目", "检查内容", "检查依据", "检查方法", "判定标准",
        "重点项", "否决项", "分值", "适用范围", "备注"]


def _row(it):
    return [it["code"], it["category"], it["name"], it["content"], it["legal_basis"], it["method"],
            it["criteria"], "是" if it["is_key"] else "", "是" if it["is_veto"] else "",
            it["score"], it["scope"], it["note"]]


def items_xlsx(items, title="检查项目库"):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(HEAD)
    for w, col in zip([8, 18, 18, 42, 34, 10, 26, 7, 7, 6, 14, 20], "ABCDEFGHIJKL"):
        ws.column_dimensions[col].width = w
    for it in items:
        ws.append(_row(it))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def items_docx(items, title="水路运输企业检查内容清单", subtitle=""):
    doc = Document()
    doc.add_heading(title, level=0)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.runs[0].font.size = Pt(10)
    doc.add_paragraph("法规依据：《国内水路运输管理条例》（2023年第三次修订）、"
                      "《国内水路运输管理规定》（2020年修正）、《中华人民共和国安全生产法》（2021年修正）。"
                      "注：船员配备不足、船舶不适航、超载、货船载客、危险货物无适装证书等属海事管理处罚事项"
                      "（条例第38条），本清单中标注“移送海事”。")
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    hdr = ["编号", "检查项目", "检查内容", "检查依据", "检查方法", "判定标准", "备注"]
    for i, h in enumerate(hdr):
        table.rows[0].cells[i].text = h
    for it in items:
        cells = table.add_row().cells
        mark = ""
        if it["is_key"]:
            mark += "重点项 "
        if it["is_veto"]:
            mark += "否决项 "
        vals = [it["code"], it["name"] + ("（%s）" % mark.strip() if mark else ""),
                it["content"], it["legal_basis"], it["method"], it["criteria"],
                (it["note"] or "") + (" 适用：" + it["scope"] if it["scope"] else "")]
        for i, v in enumerate(vals):
            cells[i].text = str(v or "")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def ent_template_xlsx():
    """企业名录批量导入模板"""
    wb = Workbook()
    ws = wb.active
    ws.title = "企业名录"
    cols = ["企业名称", "统一社会信用代码", "经营许可证号", "经营类别", "法定代表人", "安全联系人", "联系电话", "地址", "核定经营范围"]
    ws.append(cols)
    for w, col in zip([32, 22, 20, 16, 12, 12, 16, 24, 30], "ABCDEFGHI"):
        ws.column_dimensions[col].width = w
    ws.append(["平潭XX航运有限公司", "91350128XXXXXXXXXX", "闽交水运字XXXX号", "普通货物运输",
               "张三", "李四", "13800000000", "平潭综合实验区XX路X号", "国内沿海普通货物运输"])
    ws2 = wb.create_sheet("填写说明")
    for line in ["1. 第1行为表头，请勿修改列名；从第2行开始填写。",
                 "2. 企业名称为必填；统一社会信用代码重复的行自动跳过（以信用代码去重）。",
                 "3. 经营类别：普通货物运输/危险货物运输/旅客运输/船舶管理业务/船舶代理或客货代理/其他。",
                 "4. 支持 .xlsx 格式，单次最多导入500家。"]:
        ws2.append([line])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
